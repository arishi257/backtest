from __future__ import annotations

import argparse
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from backtest.config import normalize_underlying
from backtest.replay import CsvReplayFeed, register_token_tickers
from backtest.sessions import build_backtest_sessions
from backtest_5s.__main__ import (
    DEFAULT_INDEX_ZIP,
    DynamicGammaThresholdFrozenIvState,
    DynamicGammaThresholdPortfolioState,
    default_options_zip,
    load_fyers_dataset,
    parse_date_key,
)
from fit_sensex.ui.app import total_row_numeric_values, weighted_price_total


PROJECT_ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = PROJECT_ROOT / "hols.xlsx"
THRESHOLD_RATIO = 0.40
MIN_HEDGE_THRESHOLD_LOTS = 1.2


TIME_SLICE_COLUMNS = [
    "timestamp",
    "universal_mid",
    "universal_spot",
    "atm_vol",
    "time_to_expiry",
    "full_days",
    "fraction_days",
    "live_total_pnl",
    "live_options_pv",
    "live_delta_pct",
    "live_delta_ccy",
    "live_delta_lots",
    "live_gamma_ccy_10bps",
    "live_gamma_lots_10bps",
    "live_vega_ccy_10bps",
    "live_theta_ccy",
    "live_std_1w_vega",
    "live_hedge_lots",
    "live_combined_delta_lots",
    "live_dynamic_threshold_lots",
    "frozen_total_pnl",
    "frozen_options_pv",
    "frozen_delta_pct",
    "frozen_delta_ccy",
    "frozen_delta_lots",
    "frozen_gamma_ccy_10bps",
    "frozen_gamma_lots_10bps",
    "frozen_vega_ccy_10bps",
    "frozen_theta_ccy",
    "frozen_std_1w_vega",
    "frozen_hedge_lots",
    "frozen_combined_delta_lots",
    "frozen_dynamic_threshold_lots",
]


PORTFOLIO_COLUMNS = [
    "book",
    "lots",
    "underlying",
    "maturity",
    "strike",
    "option_type",
    "qty",
    "mult",
]


def main() -> None:
    global THRESHOLD_RATIO

    parser = argparse.ArgumentParser(description="Export FYERS 5s portfolio Greeks workbook.")
    parser.add_argument("--date", default="27052026")
    parser.add_argument(
        "--underlying",
        default="SENSEX",
        type=normalize_underlying,
        choices=("SENSEX", "NIFTY"),
    )
    parser.add_argument("--options-zip", type=Path, default=None)
    parser.add_argument("--index-zip", type=Path, default=DEFAULT_INDEX_ZIP)
    parser.add_argument("--workbook", type=Path, default=WORKBOOK_PATH)
    parser.add_argument("--dynamic-gamma-threshold-ratio", type=float, default=0.40)
    parser.add_argument("--output", type=Path, default=None)
    args = parser.parse_args()

    trade_date = parse_date_key(args.date)
    THRESHOLD_RATIO = args.dynamic_gamma_threshold_ratio
    options_zip = args.options_zip or default_options_zip(args.underlying, trade_date)
    output_path = args.output or (
        PROJECT_ROOT
        / "outputs"
        / f"{args.underlying.lower()}_5s_{trade_date:%Y%m%d}_portfolio_greeks.xlsx"
    )

    dataset = load_fyers_dataset(options_zip, args.index_zip, trade_date, args.underlying)
    replay = None
    sessions = build_backtest_sessions(
        dataset,
        args.workbook,
        lambda: replay.now(),
        refresh_ms=0,
    )
    replay = CsvReplayFeed(dataset, sessions)
    register_token_tickers(sessions, dataset)
    session = sessions[0]
    live_state = DynamicGammaThresholdPortfolioState(session=session)
    frozen_state = DynamicGammaThresholdFrozenIvState(session)

    rows = []
    portfolio_rows = []
    portfolio_captured = False

    while replay.advance():
        timestamp = replay.now()
        result = session.analytics.calculate(session.store.snapshot())
        if result is None:
            continue

        live_metrics = live_state.update(
            result,
            timestamp,
            session.config.market.funding_rate,
            session.config.market.brokerage_rate,
            THRESHOLD_RATIO,
        )
        frozen_metrics = frozen_state.update(
            result,
            timestamp,
            session.config.market.funding_rate,
            session.config.market.brokerage_rate,
            THRESHOLD_RATIO,
        )

        live_risk_rows = (
            live_state.portfolio.calculate(result) if live_state.portfolio else []
        )
        frozen_risk_rows = (
            frozen_state._frozen_risk_rows(
                result,
                session.config.market.funding_rate,
            )
            if frozen_state.portfolio
            else []
        )
        live_totals = total_row_numeric_values(live_risk_rows, result) if live_risk_rows else None
        frozen_totals = (
            total_row_numeric_values(frozen_risk_rows, result) if frozen_risk_rows else None
        )

        if live_state.portfolio and not portfolio_captured:
            portfolio_rows = [
                [
                    position.book,
                    position.lots,
                    position.underlying,
                    position.maturity,
                    position.strike,
                    position.option_type,
                    position.qty,
                    position.mult,
                ]
                for position in live_state.portfolio.positions
            ]
            portfolio_captured = True

        rows.append(
            [
                timestamp.replace(tzinfo=None),
                result.universal_mid,
                result.universal_spot,
                result.atm_vol,
                result.time,
                result.full_days,
                result.fraction_days,
                live_metrics.total_pnl,
                weighted_price_total(live_risk_rows, "mid_mkt") if live_risk_rows else None,
                total_value(live_totals, 17),
                total_value(live_totals, 18),
                total_value(live_totals, 19),
                total_value(live_totals, 20),
                total_value(live_totals, 21),
                total_value(live_totals, 22),
                total_value(live_totals, 23),
                total_value(live_totals, 24),
                live_state.cumulative_hedge_lots,
                combined_delta(total_value(live_totals, 19), live_state.cumulative_hedge_lots),
                threshold(total_value(live_totals, 21)),
                frozen_metrics.total_pnl,
                weighted_price_total(frozen_risk_rows, "mid_mkt") if frozen_risk_rows else None,
                total_value(frozen_totals, 17),
                total_value(frozen_totals, 18),
                total_value(frozen_totals, 19),
                total_value(frozen_totals, 20),
                total_value(frozen_totals, 21),
                total_value(frozen_totals, 22),
                total_value(frozen_totals, 23),
                total_value(frozen_totals, 24),
                frozen_state.cumulative_hedge_lots,
                combined_delta(
                    total_value(frozen_totals, 19),
                    frozen_state.cumulative_hedge_lots,
                ),
                threshold(total_value(frozen_totals, 21)),
            ]
        )

    workbook = build_workbook(rows, portfolio_rows, args.underlying, trade_date)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    print(output_path)
    print(len(rows))


def total_value(totals: tuple | None, index: int) -> float | None:
    if totals is None or index >= len(totals):
        return None
    value = totals[index]
    return float(value) if isinstance(value, (int, float)) else None


def combined_delta(delta_lots: float | None, hedge_lots: float) -> float | None:
    if delta_lots is None:
        return None
    return delta_lots + hedge_lots


def threshold(gamma_lots: float | None) -> float | None:
    if gamma_lots is None:
        return MIN_HEDGE_THRESHOLD_LOTS
    return max(abs(gamma_lots) * THRESHOLD_RATIO, MIN_HEDGE_THRESHOLD_LOTS)


def build_workbook(
    rows: list[list],
    portfolio_rows: list[list],
    underlying: str,
    trade_date: date,
) -> Workbook:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    slices = workbook.create_sheet("Time Slices")
    portfolio = workbook.create_sheet("Portfolio")

    slices.append(TIME_SLICE_COLUMNS)
    for row in rows:
        slices.append(row)
    format_sheet_as_table(slices, f"{underlying.title()}5sGreeks")
    slices.freeze_panes = "A2"
    slices.auto_filter.ref = slices.dimensions
    set_number_formats(slices)

    portfolio.append(PORTFOLIO_COLUMNS)
    for row in portfolio_rows:
        portfolio.append(row)
    format_sheet_as_table(portfolio, f"{underlying.title()}5sPortfolio")
    portfolio.freeze_panes = "A2"

    final = rows[-1] if rows else []
    summary_rows = [
        (f"{underlying} 5s Portfolio Greeks", None),
        ("Date", trade_date.isoformat()),
        ("Threshold Ratio", THRESHOLD_RATIO),
        ("Rows", len(rows)),
        ("Live Total PnL", value_by_column(final, "live_total_pnl")),
        ("Frozen Total PnL", value_by_column(final, "frozen_total_pnl")),
        ("Live Gamma Lots 10bps", value_by_column(final, "live_gamma_lots_10bps")),
        ("Frozen Gamma Lots 10bps", value_by_column(final, "frozen_gamma_lots_10bps")),
        ("Live Delta Lots", value_by_column(final, "live_delta_lots")),
        ("Frozen Delta Lots", value_by_column(final, "frozen_delta_lots")),
    ]
    for row in summary_rows:
        summary.append(row)
    summary["A1"].font = Font(size=14, bold=True, color="1F4E78")
    summary.merge_cells("A1:B1")
    for cell in summary["A"]:
        cell.font = Font(bold=True)
    for row_index in range(3, summary.max_row + 1):
        summary[f"B{row_index}"].number_format = "#,##0.00"
    summary["B3"].number_format = "0.00%"
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 18

    autosize(slices)
    autosize(portfolio)
    return workbook


def value_by_column(row: list, column: str):
    if not row:
        return None
    return row[TIME_SLICE_COLUMNS.index(column)]


def format_sheet_as_table(sheet, table_name: str) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    if sheet.max_row > 1:
        table = Table(displayName=table_name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)


def set_number_formats(sheet) -> None:
    percent_columns = {"atm_vol"}
    integer_columns = {"live_theta_ccy", "live_std_1w_vega", "frozen_theta_ccy", "frozen_std_1w_vega"}
    for col_index, header in enumerate(TIME_SLICE_COLUMNS, start=1):
        letter = get_column_letter(col_index)
        for row_index in range(2, sheet.max_row + 1):
            cell = sheet[f"{letter}{row_index}"]
            if header == "timestamp":
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            elif header in percent_columns:
                cell.number_format = "0.00%"
            elif header in integer_columns:
                cell.number_format = "#,##0"
            else:
                cell.number_format = "#,##0.00"


def autosize(sheet) -> None:
    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
        sheet.column_dimensions[letter].width = min(max(max_length + 2, 11), 26)


if __name__ == "__main__":
    main()
