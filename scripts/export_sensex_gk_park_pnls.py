from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_CSV = PROJECT_ROOT / "processed_data" / "20260430_SENSEX.csv"
OUTPUT_PATH = PROJECT_ROOT / "outputs" / "sensex_20260430_gk_park_pnls.xlsx"

OUTPUT_COLUMNS = [
    ("timestamp", "Timestamp"),
    ("trade_date", "Trade Date"),
    ("underlying", "Underlying"),
    ("expiry", "Expiry"),
    ("portfolio_gamma_l", "Gamma (L)"),
    ("future_prev_close", "Prev Close"),
    ("future_open", "Open"),
    ("future_high", "High"),
    ("future_low", "Low"),
    ("future_close", "Close"),
    ("park_move", "Park Move"),
    ("park_gamma_pnl", "Park Gamma PnL"),
    ("park_c2c_gamma_pnl", "C2C Gamma PnL"),
    ("park_gamma_pnl_diff", "Park Gamma PnL Diff"),
    ("park_gamma_pnl_diff_total", "Park Gamma PnL Diff Total"),
    ("gk_move", "GK Move"),
    ("gk_gamma_pnl", "GK Gamma PnL"),
    ("gk_gamma_pnl_diff", "GK Gamma PnL Diff"),
    ("gk_gamma_pnl_diff_total", "GK Gamma PnL Diff Total"),
    ("portfolio_gamma_diff_total", "Gamma Diff Total"),
]


def coerce_value(value: str) -> object:
    if value == "":
        return None
    try:
        return float(value)
    except ValueError:
        return value


def load_rows() -> list[dict[str, object]]:
    if not SOURCE_CSV.exists():
        raise FileNotFoundError(f"Source CSV not found: {SOURCE_CSV}")

    with SOURCE_CSV.open(newline="") as file:
        reader = csv.DictReader(file)
        return [
            {label: coerce_value(row[key]) for key, label in OUTPUT_COLUMNS}
            for row in reader
        ]


def autosize_columns(sheet) -> None:
    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
        sheet.column_dimensions[letter].width = min(max(max_length + 2, 11), 24)


def build_workbook(rows: list[dict[str, object]]) -> Workbook:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    data = workbook.create_sheet("Slice PnLs")

    headers = [label for _, label in OUTPUT_COLUMNS]
    data.append(headers)
    for row in rows:
        data.append([row[label] for label in headers])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in data[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    data.freeze_panes = "A2"
    data.auto_filter.ref = data.dimensions

    table = Table(displayName="SensexPnLSlices", ref=data.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    data.add_table(table)

    pct_columns = {"Park Move", "GK Move"}
    pnl_columns = {
        "Gamma (L)",
        "Park Gamma PnL",
        "C2C Gamma PnL",
        "Park Gamma PnL Diff",
        "Park Gamma PnL Diff Total",
        "GK Gamma PnL",
        "GK Gamma PnL Diff",
        "GK Gamma PnL Diff Total",
        "Gamma Diff Total",
    }
    price_columns = {"Prev Close", "Open", "High", "Low", "Close"}
    for col_idx, header in enumerate(headers, start=1):
        letter = get_column_letter(col_idx)
        if header in pct_columns:
            for cell in data[f"{letter}2:{letter}{data.max_row}"]:
                cell[0].number_format = "0.0000%"
        elif header in pnl_columns:
            for cell in data[f"{letter}2:{letter}{data.max_row}"]:
                cell[0].number_format = "#,##0.00"
        elif header in price_columns:
            for cell in data[f"{letter}2:{letter}{data.max_row}"]:
                cell[0].number_format = "#,##0.00"

    for header in ("Park Gamma PnL Diff", "GK Gamma PnL Diff"):
        col_idx = headers.index(header) + 1
        letter = get_column_letter(col_idx)
        data.conditional_formatting.add(
            f"{letter}2:{letter}{data.max_row}",
            ColorScaleRule(
                start_type="min",
                start_color="F8696B",
                mid_type="percentile",
                mid_value=50,
                mid_color="FFEB84",
                end_type="max",
                end_color="63BE7B",
            ),
        )

    autosize_columns(data)
    data.column_dimensions["A"].width = 21

    final = rows[-1] if rows else {}
    summary_rows = [
        ("SENSEX 0DTE GK / Parkinson Gamma PnLs", None),
        ("Date", "2026-04-30"),
        ("Slices", len(rows)),
        ("Final Park Gamma PnL Diff Total", final.get("Park Gamma PnL Diff Total")),
        ("Final GK Gamma PnL Diff Total", final.get("GK Gamma PnL Diff Total")),
        ("Final Gamma Diff Total", final.get("Gamma Diff Total")),
        ("Final Gamma (L)", final.get("Gamma (L)")),
    ]
    for row in summary_rows:
        summary.append(row)

    summary["A1"].font = Font(size=14, bold=True, color="1F4E78")
    summary.merge_cells("A1:B1")
    for cell in summary["A"]:
        cell.font = Font(bold=True)
    for row in range(3, summary.max_row + 1):
        summary[f"B{row}"].number_format = "#,##0.00"
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 18

    return workbook


def main() -> None:
    rows = load_rows()
    workbook = build_workbook(rows)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT_PATH)
    print(OUTPUT_PATH)
    print(len(rows))


if __name__ == "__main__":
    main()
