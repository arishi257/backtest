from __future__ import annotations

import argparse
import csv
import os
import shutil
import subprocess
import sys
import time
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
SCRIPTS = ROOT / "scripts"
for path in (SRC, SCRIPTS):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from backtest.config import DEFAULT_WORKBOOK, normalize_underlying
from backtest.data import OptionDataset
from backtest.replay import CsvReplayFeed, register_token_tickers
from backtest.sessions import build_backtest_sessions
from backtest_1m.__main__ import dynamic_gamma_threshold
from backtest_sample_hf.__main__ import parse_date_key
from export_sample_hf_diagnostics import (
    close_to_close_variance,
    round_or_none,
    scaled_volatility_or_none,
)
from run_sample_hf_pnl_only import (
    DirectHedgeState,
    DirectPriceBook,
    capture_light_frozen_ivs,
)

NODE_EXE = Path(
    r"C:\Users\rishi\.cache\codex-runtimes\codex-primary-runtime\dependencies\node\bin\node.exe"
)
IST = ZoneInfo("Asia/Kolkata")
PNL_START_TIME = "09:20"


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Run light diagnostics from a Breeze 1-second workbook using a saved portfolio."
    )
    parser.add_argument("--date", required=True)
    parser.add_argument("--underlying", required=True, type=normalize_underlying)
    parser.add_argument("--breeze-workbook", type=Path, required=True)
    parser.add_argument("--portfolio-workbook", type=Path, required=True)
    parser.add_argument(
        "--replace-position",
        action="append",
        default=[],
        help="Replace a portfolio leg as OLD_STRIKE:TYPE:NEW_STRIKE, e.g. 23750:CE:23800.",
    )
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--dynamic-gamma-threshold-ratio", type=float, default=0.40)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "runs_1s" / "Light_Run")
    args = parser.parse_args()

    trade_date = parse_date_key(args.date)
    run_dir = ROOT / "runs_1s" / f"_breeze_light_tmp_{args.underlying.lower()}_{trade_date:%Y%m%d}_{datetime.now():%Y%m%d_%H%M%S}"
    run_dir.mkdir(parents=True, exist_ok=True)

    total_start = time.perf_counter()
    data_start = time.perf_counter()
    dataset = load_breeze_dataset(args.breeze_workbook, trade_date, args.underlying)
    positions = load_portfolio_positions(args.portfolio_workbook)
    apply_position_replacements(positions, args.replace_position)
    data_seconds = time.perf_counter() - data_start

    sim_start = time.perf_counter()
    diagnostics, portfolio_rows, summary, cycles = run_breeze_light_diagnostics(
        dataset,
        positions,
        args.workbook,
        args.dynamic_gamma_threshold_ratio,
    )
    simulation_seconds = time.perf_counter() - sim_start

    write_csv(run_dir / "diagnostics.csv", diagnostics)
    write_csv(run_dir / "portfolio_0920.csv", portfolio_rows)
    write_summary_csv(run_dir / "summary_metrics.csv", summary)
    workbook_start = time.perf_counter()
    temp_workbook = build_diagnostics_workbook(run_dir, args.underlying, trade_date)
    workbook_seconds = time.perf_counter() - workbook_start

    args.output_dir.mkdir(parents=True, exist_ok=True)
    final_workbook = unique_output_path(
        args.output_dir,
        args.underlying.lower(),
        trade_date,
    )
    shutil.move(str(temp_workbook), final_workbook)
    shutil.rmtree(run_dir, ignore_errors=True)
    total_seconds = time.perf_counter() - total_start

    print(f"Completed {cycles} Breeze light 1-second cycles.")
    print(f"Workbook: {final_workbook}")
    print("Final metrics:")
    for key in (
        "portfolio_total_pnl",
        "portfolio_gamma_l",
        "portfolio_gamma_diff_total",
        "frozen_iv_total_pnl",
        "c2c_synth_vol",
        "hedge_vol",
    ):
        print(f"  {key}: {format_number(summary.get(key))}")
    print("Timing:")
    print(f"  data_load_seconds: {data_seconds:.2f}")
    print(f"  simulation_seconds: {simulation_seconds:.2f}")
    print(f"  workbook_seconds: {workbook_seconds:.2f}")
    print(f"  total_seconds: {total_seconds:.2f}")


def load_breeze_dataset(path: Path, trade_date: date, underlying: str) -> OptionDataset:
    columns = [
        "close",
        "datetime",
        "expiry_date",
        "product_type",
        "right",
        "stock_code",
        "strike_price",
        "trade_date",
        "underlying",
    ]
    if path.suffix.lower() == ".parquet":
        raw = pd.read_parquet(path, columns=columns)
    else:
        raw = pd.read_excel(path, sheet_name="data", usecols=columns)
    raw = raw[
        raw["underlying"].astype(str).str.upper().eq(underlying)
        & raw["product_type"].astype(str).str.lower().eq("options")
    ].copy()
    if raw.empty:
        raise ValueError(f"No {underlying} option rows found in {path}.")
    raw["timestamp"] = pd.to_datetime(raw["datetime"], errors="coerce")
    raw["timestamp"] = raw["timestamp"].dt.tz_localize(IST)
    raw["expiry"] = pd.to_datetime(raw["expiry_date"], errors="coerce").dt.date
    raw["strike"] = pd.to_numeric(raw["strike_price"], errors="coerce").astype("Int64")
    raw["option_type"] = raw["right"].map({"call": "CE", "put": "PE", "Call": "CE", "Put": "PE"})
    raw["close"] = pd.to_numeric(raw["close"], errors="coerce")
    raw = raw.dropna(subset=["timestamp", "expiry", "strike", "option_type", "close"])
    raw = raw[raw["timestamp"].dt.date.eq(trade_date)]
    raw = raw[raw["expiry"].eq(trade_date)]
    if raw.empty:
        raise ValueError(f"No {underlying} rows found for {trade_date} in {path}.")
    raw["strike"] = raw["strike"].astype(int)
    raw["ticker"] = raw.apply(
        lambda row: f"{underlying}{row['expiry']:%d%b%y}{int(row['strike'])}{row['option_type']}.NFO".upper(),
        axis=1,
    )
    frame = raw[["ticker", "timestamp", "expiry", "strike", "option_type", "close"]].sort_values(
        ["timestamp", "ticker"]
    )
    available = pd.DatetimeIndex(frame["timestamp"].drop_duplicates().sort_values())
    start = pd.Timestamp(f"{trade_date} {PNL_START_TIME}:00", tz=IST)
    end = pd.Timestamp(f"{trade_date} 15:24:59", tz=IST)
    available = available[(available >= start) & (available <= end)]
    if available.empty:
        raise ValueError(f"No Breeze timestamps found between {start} and {end}.")
    timestamps = pd.date_range(available.min(), available.max(), freq="s")
    return OptionDataset(frame.reset_index(drop=True), trade_date, timestamps, underlying, future_series=None)


def load_portfolio_positions(path: Path) -> list[SimpleNamespace]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Summary"]
    header_index = None
    headers = []
    for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        values = [str(value).strip() if value is not None else "" for value in row]
        if values[:9] == [
            "timestamp",
            "universal_mid",
            "underlying",
            "maturity",
            "strike",
            "option_type",
            "lots",
            "qty",
            "mult",
        ]:
            header_index = index
            headers = values
            break
    if header_index is None:
        raise ValueError(f"Could not find portfolio table in {path}.")
    positions = []
    for row in sheet.iter_rows(min_row=header_index + 1, values_only=True):
        if row[0] is None:
            continue
        record = dict(zip(headers, row))
        positions.append(
            SimpleNamespace(
                underlying=record["underlying"],
                maturity=record["maturity"],
                strike=int(record["strike"]),
                option_type=str(record["option_type"]).upper(),
                lots=float(record["lots"]),
                qty=float(record["qty"]),
                mult=float(record["mult"]),
            )
        )
    workbook.close()
    if not positions:
        raise ValueError(f"No portfolio positions found in {path}.")
    return positions


def apply_position_replacements(positions: list[SimpleNamespace], replacements: list[str]) -> None:
    for replacement in replacements:
        try:
            old_strike_text, option_type, new_strike_text = replacement.split(":")
            old_strike = int(old_strike_text)
            new_strike = int(new_strike_text)
            normalized_type = option_type.upper()
        except ValueError as exc:
            raise ValueError(
                f"Invalid --replace-position {replacement!r}; expected OLD_STRIKE:TYPE:NEW_STRIKE."
            ) from exc
        matched = False
        for position in positions:
            if position.strike == old_strike and position.option_type == normalized_type:
                position.strike = new_strike
                matched = True
        if not matched:
            raise ValueError(f"No portfolio leg matched replacement {replacement!r}.")
    merge_duplicate_positions(positions)


def merge_duplicate_positions(positions: list[SimpleNamespace]) -> None:
    merged: dict[tuple[str, str, int, str], SimpleNamespace] = {}
    for position in positions:
        key = (position.underlying, position.maturity, position.strike, position.option_type)
        existing = merged.get(key)
        if existing is None:
            merged[key] = position
        else:
            existing.lots += position.lots
            existing.qty += position.qty
    positions[:] = sorted(
        merged.values(),
        key=lambda position: (position.strike, position.option_type),
    )


def run_breeze_light_diagnostics(dataset, positions, workbook: Path, threshold_ratio: float):
    replay = None
    current_time = lambda: replay.now()
    sessions = build_backtest_sessions(dataset, workbook, current_time, refresh_ms=0)
    replay = CsvReplayFeed(dataset, sessions)
    register_token_tickers(sessions, dataset)
    if len(sessions) != 1:
        raise ValueError("Expected one expiry session.")
    session = sessions[0]
    surface = dataset.option_surface(session.spec.expiry)
    price_book = DirectPriceBook(dataset, surface, session)
    portfolio = SimpleNamespace(positions=positions)

    market_state = None
    frozen_state = None
    diagnostics = []
    portfolio_rows = []
    universal_mid_points = []
    gamma_l_points = []
    previous_universal_mid = None
    previous_hedge_universal_mid = None
    c2c_synth_variance = 0.0
    hedge_variance = 0.0
    calendar_days = session.config.market.calendar_days
    intraday_var = session.config.market.intraday_var
    user_value = session.config.market.user_value
    summary = {}
    cycles = 0

    while replay.advance():
        timestamp = replay.now()
        if timestamp.strftime("%H:%M") < PNL_START_TIME:
            continue
        cycles += 1
        timestamp_value = dataset.timestamps[replay.position]
        row = surface.loc[timestamp_value]
        snapshot = price_book.market_snapshot(row, timestamp, user_value)
        if snapshot is None:
            continue
        user_value = snapshot.user_value
        if market_state is None or frozen_state is None:
            complete = complete_strikes(price_book, row)
            if not complete:
                continue
            hedge_strike = min(complete, key=lambda strike: (abs(strike - snapshot.universal_mid), strike))
            frozen_ivs = capture_light_frozen_ivs(positions, price_book, row, snapshot)
            if len(frozen_ivs) != len(positions):
                continue
            market_state = DirectHedgeState(portfolio=portfolio, options_pv_snapshot=None, hedge_strike=hedge_strike)
            frozen_state = DirectHedgeState(
                portfolio=portfolio,
                options_pv_snapshot=None,
                hedge_strike=hedge_strike,
                frozen_ivs=frozen_ivs,
            )
            portfolio_rows = portfolio_snapshot_rows(timestamp, snapshot.universal_mid, positions)

        prior_trade_count = len(market_state.hedge_trades)
        prior_frozen_trade_count = len(frozen_state.hedge_trades)
        market_pnl = market_state.update(price_book, row, snapshot, timestamp, threshold_ratio)
        frozen_pnl = frozen_state.update(price_book, row, snapshot, timestamp, threshold_ratio)
        if market_pnl is None:
            continue

        gamma_lots = market_state.last_gamma_lots
        threshold_lots = dynamic_gamma_threshold(gamma_lots, threshold_ratio)
        net_delta_lots = market_state.last_delta_lots + market_state.cumulative_hedge_lots
        traded_delta_lots = sum(float(trade["lots_change"]) for trade in market_state.hedge_trades[prior_trade_count:])
        traded_universal_mid = snapshot.universal_mid if abs(traded_delta_lots) > 1e-9 else None

        frozen_threshold_lots = dynamic_gamma_threshold(frozen_state.last_gamma_lots, threshold_ratio)
        frozen_net_delta_lots = frozen_state.last_delta_lots + frozen_state.cumulative_hedge_lots
        frozen_traded_delta_lots = sum(
            float(trade["lots_change"]) for trade in frozen_state.hedge_trades[prior_frozen_trade_count:]
        )
        frozen_traded_universal_mid = snapshot.universal_mid if abs(frozen_traded_delta_lots) > 1e-9 else None

        if previous_universal_mid is not None and timestamp.strftime("%H:%M:%S").endswith(":00"):
            c2c_synth_variance += close_to_close_variance(previous_universal_mid, snapshot.universal_mid)
        if timestamp.strftime("%H:%M:%S").endswith(":00"):
            previous_universal_mid = snapshot.universal_mid
        if traded_universal_mid is not None:
            if previous_hedge_universal_mid is not None:
                hedge_variance += close_to_close_variance(previous_hedge_universal_mid, traded_universal_mid)
            previous_hedge_universal_mid = traded_universal_mid

        if gamma_lots is not None:
            gamma_l_points.append((timestamp, gamma_lots))
        universal_mid_points.append((timestamp, snapshot.universal_mid))

        diagnostics.append(
            {
                "timestamp": timestamp.strftime("%Y-%m-%d %H:%M:%S"),
                "universal_mid": round_or_none(snapshot.universal_mid),
                "running_total_pnl": round_or_none(market_pnl),
                "gamma_lots": round_or_none(gamma_lots),
                "threshold_lots": round_or_none(threshold_lots),
                "net_delta_lots_options_plus_hedge": round_or_none(net_delta_lots),
                "hedge_delta_lots": round_or_none(market_state.cumulative_hedge_lots),
                "traded_delta_lots": round_or_none(traded_delta_lots),
                "traded_universal_mid": round_or_none(traded_universal_mid),
                "frozen_iv_running_total_pnl": round_or_none(frozen_pnl),
                "frozen_iv_gamma_lots": round_or_none(frozen_state.last_gamma_lots),
                "frozen_iv_threshold_lots": round_or_none(frozen_threshold_lots),
                "frozen_iv_net_delta_lots_options_plus_hedge": round_or_none(frozen_net_delta_lots),
                "frozen_iv_hedge_delta_lots": round_or_none(frozen_state.cumulative_hedge_lots),
                "frozen_iv_traded_delta_lots": round_or_none(frozen_traded_delta_lots),
                "frozen_iv_traded_universal_mid": round_or_none(frozen_traded_universal_mid),
            }
        )
        summary = {
            "portfolio_total_pnl": market_pnl,
            "portfolio_gamma_l": gamma_lots,
            "portfolio_gamma_diff_total": None,
            "park_gamma_pnl_diff_total": None,
            "gk_gamma_pnl_diff_total": None,
            "frozen_iv_total_pnl": frozen_pnl,
            "c2c_spot_vol": None,
            "park_vol": None,
            "gk_vol": None,
            "c2c_synth_vol": scaled_volatility_or_none(c2c_synth_variance, calendar_days, intraday_var),
            "hedge_vol": scaled_volatility_or_none(hedge_variance, calendar_days, intraday_var),
        }

    if not diagnostics:
        raise ValueError("No Breeze light diagnostics rows were produced.")
    return diagnostics, portfolio_rows, summary, cycles


def complete_strikes(price_book: DirectPriceBook, row) -> list[int]:
    return [
        strike
        for strike in price_book.strikes
        if price_book.option_price(row, strike, "CE") is not None
        and price_book.option_price(row, strike, "PE") is not None
    ]


def portfolio_snapshot_rows(timestamp, universal_mid, positions) -> list[dict[str, object]]:
    return [
        {
            "timestamp": timestamp.strftime("%Y-%m-%d %H:%M:%S"),
            "universal_mid": round_or_none(universal_mid),
            "underlying": position.underlying,
            "maturity": position.maturity,
            "strike": position.strike,
            "option_type": position.option_type,
            "lots": position.lots,
            "qty": position.qty,
            "mult": position.mult,
        }
        for position in positions
    ]


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def write_summary_csv(path: Path, summary: dict[str, object]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=["metric", "value"])
        writer.writeheader()
        for key in (
            "portfolio_total_pnl",
            "portfolio_gamma_l",
            "portfolio_gamma_diff_total",
            "park_gamma_pnl_diff_total",
            "gk_gamma_pnl_diff_total",
            "frozen_iv_total_pnl",
            "c2c_spot_vol",
            "park_vol",
            "gk_vol",
            "c2c_synth_vol",
            "hedge_vol",
        ):
            writer.writerow({"metric": key, "value": round_or_none(summary.get(key))})


def build_diagnostics_workbook(run_dir: Path, underlying: str, trade_date: date) -> Path:
    env = os.environ.copy()
    env["RUN_DIR"] = str(run_dir)
    env["UNDERLYING"] = underlying
    env["TRADE_DATE"] = trade_date.isoformat()
    env["CLEAN_INPUTS"] = "1"
    env["NODE_PATH"] = r"C:\Users\rishi\.cache\codex-runtimes\codex-primary-runtime\dependencies\node\node_modules"
    subprocess.run(
        [str(NODE_EXE), str(ROOT / "scripts" / "build_sample_hf_diagnostics_workbook.mjs")],
        cwd=ROOT,
        env=env,
        check=True,
    )
    workbooks = sorted(run_dir.glob("*.xlsx"))
    if not workbooks:
        raise SystemExit(f"No workbook was created in {run_dir}.")
    return workbooks[-1]


def unique_output_path(output_dir: Path, underlying: str, trade_date: date) -> Path:
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    date_name = f"{trade_date.day:02d}{months[trade_date.month - 1]}{trade_date.year % 100:02d}"
    base = f"{underlying}_{date_name}_light_diagnostics"
    path = output_dir / f"{base}.xlsx"
    suffix = 2
    while path.exists():
        path = output_dir / f"{base}_{suffix}.xlsx"
        suffix += 1
    return path


def format_number(value: object) -> str:
    if value is None:
        return "--"
    try:
        return f"{float(value):.2f}"
    except (TypeError, ValueError):
        return str(value)


if __name__ == "__main__":
    main()
