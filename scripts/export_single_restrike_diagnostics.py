from __future__ import annotations

import argparse
import sys
import time
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
SCRIPTS = ROOT / "scripts"
for path in (SRC, SCRIPTS):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from backtest.config import DEFAULT_WORKBOOK, normalize_underlying
from backtest_sample_hf.__main__ import load_sample_hf_option_dataset, parse_date_key
from export_1s_batch_restrike_workbooks import run_path, write_path_sheet


DEFAULT_DATA_DIR = Path(r"C:\options data\1s data")
DEFAULT_OUTPUT_DIR = ROOT / "10am_restrike_anurag"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--date", required=True)
    parser.add_argument("--underlying", required=True, type=normalize_underlying, choices=("NIFTY", "SENSEX"))
    parser.add_argument("--data-dir", type=Path, default=DEFAULT_DATA_DIR)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--start-time", default="10:00")
    parser.add_argument("--end-time", default="15:15")
    parser.add_argument("--threshold-ratio", type=float, default=0.10)
    args = parser.parse_args()

    trade_date = parse_date_key(args.date)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    threshold_label = f"{args.threshold_ratio:.2f}".replace(".", "p").rstrip("0").rstrip("p")
    output = args.output_dir / (
        f"{args.underlying.lower()}_{trade_date:%d%b%y}_"
        f"{args.start_time.replace(':', '')}_threshold_{threshold_label}_restrike_diagnostics.xlsx"
    )

    total_start = time.perf_counter()
    load_start = time.perf_counter()
    dataset = load_sample_hf_option_dataset(args.data_dir, trade_date, args.underlying)
    load_seconds = time.perf_counter() - load_start
    sim_start = time.perf_counter()
    result = run_path(
        dataset,
        DEFAULT_WORKBOOK,
        args.threshold_ratio,
        args.start_time,
        args.end_time,
        with_restrike=True,
    )
    sim_seconds = time.perf_counter() - sim_start
    total_seconds = time.perf_counter() - total_start
    write_workbook(
        output,
        args.underlying,
        trade_date,
        args.start_time,
        args.end_time,
        args.threshold_ratio,
        result,
        load_seconds,
        sim_seconds,
        total_seconds,
    )
    summary = result["summary"]
    print(f"OUTPUT,{output}")
    print(f"FINAL_PNL,{summary.get('final_pnl'):.2f}")


def write_workbook(
    output: Path,
    underlying: str,
    trade_date,
    start_time: str,
    end_time: str,
    threshold_ratio: float,
    result,
    load_seconds: float,
    sim_seconds: float,
    total_seconds: float,
) -> None:
    summary = result["summary"]
    gammas = [row["gamma_lots"] for row in result["rows"] if row.get("gamma_lots") is not None]
    avg_gamma = sum(gammas) / len(gammas) if gammas else None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"{underlying} {trade_date:%d-%b-%y} 10am Restrike Diagnostics"
    ws["A1"].font = Font(bold=True, size=14)
    rows = [
        ("Underlying", underlying),
        ("Trade Date", trade_date.strftime("%d-%b-%y")),
        ("Start Time", start_time),
        ("End Time", end_time),
        ("Threshold Ratio", threshold_ratio),
        ("Restrike", "Yes"),
        ("Final PnL", summary.get("final_pnl")),
        ("c2c_vol", summary.get("c2c_vol")),
        ("hedge_vol", summary.get("hedge_vol")),
        ("PnL at Breach", summary.get("pnl_at_breach")),
        ("Breach Timestamp", summary.get("breach_timestamp")),
        ("Breach Side", summary.get("breach_side")),
        ("Breach Strike", summary.get("breach_strike")),
        ("Max Net Delta Lots", summary.get("max_net_delta_lots")),
        ("Min Net Delta Lots", summary.get("min_net_delta_lots")),
        ("Avg Net Delta Lots", summary.get("avg_net_delta_lots")),
        ("Average Gamma Lots", avg_gamma),
        ("Number of Hedge Events", summary.get("synthetic_hedge_trades")),
        ("Synthetic Buy Lots Hedged", summary.get("synthetic_buy_lots")),
        ("Synthetic Sell Lots Hedged", summary.get("synthetic_sell_lots")),
        ("Total Synthetic Lots Hedged", summary.get("synthetic_total_abs_lots")),
        ("Total Premium Bought", summary.get("premium_bought")),
        ("Total Premium Sold", summary.get("premium_sold")),
        ("Brokerage", brokerage(summary.get("premium_bought"), summary.get("premium_sold"))),
        ("Slippage", slippage(underlying, summary.get("synthetic_total_abs_lots"))),
        ("Data Load Seconds", load_seconds),
        ("Simulation Seconds", sim_seconds),
        ("Total Seconds", total_seconds),
    ]
    for row_idx, (label, value) in enumerate(rows, start=3):
        ws.cell(row_idx, 1).value = label
        ws.cell(row_idx, 1).font = Font(bold=True)
        ws.cell(row_idx, 2).value = value
        if isinstance(value, (int, float)):
            ws.cell(row_idx, 2).number_format = "#,##0" if label in {"Brokerage", "Slippage"} else "#,##0.00"

    path_ws = wb.create_sheet("Restrike Path")
    write_path_sheet(path_ws, result["rows"])
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A3" if sheet.title == "Summary" else "A2"
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 18
    wb.save(output)


def brokerage(premium_bought, premium_sold):
    if premium_bought is None or premium_sold is None:
        return None
    return 0.0024 * ((float(premium_bought) + float(premium_sold)) / 2)


def slippage(underlying: str, synthetic_lots):
    if synthetic_lots is None:
        return None
    normalized = underlying.strip().upper()
    if normalized == "NIFTY":
        return float(synthetic_lots) * 0.125 * 65
    if normalized == "SENSEX":
        return float(synthetic_lots) * 0.41 * 20
    return None


if __name__ == "__main__":
    main()
