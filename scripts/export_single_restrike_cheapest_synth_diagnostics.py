from __future__ import annotations

import argparse
import sys
import time
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
SCRIPTS = ROOT / "scripts"
for path in (SRC, SCRIPTS):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from backtest.config import DEFAULT_WORKBOOK, normalize_underlying
from backtest.replay import CsvReplayFeed, register_token_tickers
from backtest.sessions import build_backtest_sessions
from backtest_1m.__main__ import dynamic_gamma_threshold
from backtest_sample_hf.__main__ import load_sample_hf_option_dataset, parse_date_key
from export_1s_batch_restrike_workbooks import (
    SegmentInfo,
    in_time_window,
    portfolio_liquidation_premiums,
    portfolio_premiums,
    synthetic_trade_premiums,
)
from export_sample_hf_diagnostics import close_to_close_variance, scaled_volatility_or_none
from run_nifty_1m_light_batch import complete_strikes
from run_sample_hf_pnl_only import (
    DirectHedgeState,
    DirectPriceBook,
    build_light_sample_portfolio,
    capture_light_frozen_ivs,
    direct_portfolio_values,
    hedge_multiplier,
)


DEFAULT_DATA_DIR = Path(r"C:\options data\1s data")
DEFAULT_OUTPUT_DIR = ROOT / "10am_restrike_anurag"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--date", required=True)
    parser.add_argument("--underlying", required=True, type=normalize_underlying, choices=("NIFTY", "SENSEX"))
    parser.add_argument("--data-dir", type=Path, default=DEFAULT_DATA_DIR)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--start-time", default="10:00")
    parser.add_argument("--end-time", default="15:15")
    parser.add_argument("--threshold-ratio", type=float, default=0.10)
    parser.add_argument(
        "--hedge-strike-rule",
        choices=(
            "cheapest",
            "cheapest_delayed",
            "closest_um",
            "regular_delayed",
            "best_market",
            "best_market_delayed",
        ),
        default="cheapest",
        help=(
            "Synthetic hedge strike rule: lowest CE+PE premium, closest strike to current UM, "
            "best executable synthetic bid/ask for the trade direction, or next-slice delayed execution."
        ),
    )
    args = parser.parse_args()

    trade_date = parse_date_key(args.date)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    threshold_label = f"{args.threshold_ratio:.2f}".replace(".", "p").rstrip("0").rstrip("p")
    rule_labels = {
        "cheapest": "cheapest_synth",
        "cheapest_delayed": "cheapest_delayed_synth",
        "closest_um": "closest_um_synth",
        "regular_delayed": "regular_delayed_synth",
        "best_market": "best_market_synth",
        "best_market_delayed": "best_market_delayed_synth",
    }
    rule_label = rule_labels[args.hedge_strike_rule]
    output = args.output_dir / (
        f"{args.underlying.lower()}_{trade_date:%d%b%y}_"
        f"{args.start_time.replace(':', '')}_threshold_{threshold_label}_"
        f"{rule_label}_restrike_diagnostics.xlsx"
    )

    total_start = time.perf_counter()
    load_start = time.perf_counter()
    dataset = load_sample_hf_option_dataset(args.data_dir, trade_date, args.underlying)
    load_seconds = time.perf_counter() - load_start
    sim_start = time.perf_counter()
    result = run_cheapest_synth_path(
        dataset,
        DEFAULT_WORKBOOK,
        args.threshold_ratio,
        args.start_time,
        args.end_time,
        args.hedge_strike_rule,
    )
    sim_seconds = time.perf_counter() - sim_start
    total_seconds = time.perf_counter() - total_start
    write_workbook(
        output,
        args.underlying,
        trade_date,
        args.start_time,
        args.end_time,
        args.threshold_ratio,
        args.hedge_strike_rule,
        result,
        load_seconds,
        sim_seconds,
        total_seconds,
    )
    print(f"OUTPUT,{output}")
    print(f"FINAL_PNL,{result['summary'].get('final_pnl'):.2f}")
    print(f"HEDGE_STRIKE_INITIAL,{result['summary'].get('initial_hedge_strike')}")


def run_cheapest_synth_path(
    dataset,
    workbook: Path,
    threshold_ratio: float,
    start_time: str,
    end_time: str,
    hedge_strike_rule: str,
):
    replay = None
    current_time = lambda: replay.now()
    sessions = build_backtest_sessions(dataset, workbook, current_time, refresh_ms=0)
    replay = CsvReplayFeed(dataset, sessions)
    register_token_tickers(sessions, dataset)
    if len(sessions) != 1:
        raise ValueError("Expected one expiry session.")
    session = sessions[0]
    surface = dataset.option_surface(session.spec.expiry)
    price_book = DirectPriceBook(dataset, surface, session)

    state_info: SegmentInfo | None = None
    waiting_reentry = False
    reentry_done = False
    user_value = session.config.market.user_value
    realized_offset = 0.0
    breach = None
    rows = []
    prev_um = None
    prev_hedge_um = None
    c2c_var = 0.0
    hedge_var = 0.0
    net_deltas = []
    hedge_trade_count = 0
    synthetic_buy_lots = 0.0
    synthetic_sell_lots = 0.0
    premium_bought = 0.0
    premium_sold = 0.0
    multiplier = None
    initial_hedge_strike = None

    while replay.advance():
        timestamp = replay.now()
        if not in_time_window(timestamp, start_time, end_time):
            continue
        row = surface.loc[dataset.timestamps[replay.position]]

        if waiting_reentry:
            state_info = make_segment(session, price_book, row, timestamp, user_value, hedge_strike_rule)
            if state_info is None:
                continue
            multiplier = hedge_multiplier(state_info.state.portfolio.positions)
            bought, sold = portfolio_premiums(state_info.state.portfolio.positions, price_book, row, multiplier)
            premium_bought += bought
            premium_sold += sold
            user_value = state_info.user_value
            waiting_reentry = False

        if state_info is None:
            state_info = make_segment(session, price_book, row, timestamp, user_value, hedge_strike_rule)
            if state_info is None:
                continue
            if initial_hedge_strike is None:
                initial_hedge_strike = state_info.hedge_strike
            multiplier = hedge_multiplier(state_info.state.portfolio.positions)
            bought, sold = portfolio_premiums(state_info.state.portfolio.positions, price_book, row, multiplier)
            premium_bought += bought
            premium_sold += sold
            user_value = state_info.user_value

        snapshot = price_book.market_snapshot(row, timestamp, user_value)
        if snapshot is None:
            continue
        user_value = snapshot.user_value
        prior_trade_count = len(state_info.state.hedge_trades)
        segment_pnl = state_info.state.update(price_book, row, snapshot, timestamp, threshold_ratio)
        if segment_pnl is None:
            continue

        new_trades = state_info.state.hedge_trades[prior_trade_count:]
        for trade in new_trades:
            lots_change = float(trade["lots_change"])
            if lots_change > 0:
                synthetic_buy_lots += lots_change
            elif lots_change < 0:
                synthetic_sell_lots += abs(lots_change)
            bought, sold = synthetic_trade_premiums(
                price_book,
                row,
                int(trade.get("hedge_strike", state_info.hedge_strike)),
                lots_change,
                multiplier,
            )
            premium_bought += bought
            premium_sold += sold
        hedge_trade_count += len(new_trades)

        running_pnl = realized_offset + segment_pnl
        net_delta = (state_info.state.last_delta_lots or 0.0) + state_info.state.cumulative_hedge_lots
        gamma_lots = state_info.state.last_gamma_lots
        threshold_lots = dynamic_gamma_threshold(gamma_lots, threshold_ratio)
        net_deltas.append(net_delta)
        if prev_um is not None:
            c2c_var += close_to_close_variance(prev_um, snapshot.universal_mid)
        prev_um = snapshot.universal_mid
        traded_um = snapshot.universal_mid if new_trades else None
        if traded_um is not None:
            if prev_hedge_um is not None:
                hedge_var += close_to_close_variance(prev_hedge_um, traded_um)
            prev_hedge_um = traded_um

        rows.append(
            {
                "timestamp": timestamp,
                "universal_mid": snapshot.universal_mid,
                "running_total_pnl": running_pnl,
                "net_delta_lots": net_delta,
                "gamma_lots": gamma_lots,
                "threshold_lots": threshold_lots,
                "hedge_delta_lots": state_info.state.cumulative_hedge_lots,
                "traded_delta_lots": sum(float(trade["lots_change"]) for trade in new_trades),
                "hedge_strike": state_info.hedge_strike,
            }
        )

        if not reentry_done:
            breach_side = None
            breach_strike = None
            if state_info.short_call is not None and snapshot.universal_mid > state_info.short_call:
                breach_side = "CALL"
                breach_strike = state_info.short_call
            elif state_info.short_put is not None and snapshot.universal_mid < state_info.short_put:
                breach_side = "PUT"
                breach_strike = state_info.short_put
            if breach_side:
                bought, sold = portfolio_liquidation_premiums(
                    state_info.state.portfolio.positions,
                    price_book,
                    row,
                    multiplier,
                )
                premium_bought += bought
                premium_sold += sold
                breach = {
                    "side": breach_side,
                    "timestamp": timestamp,
                    "strike": breach_strike,
                    "universal_mid": snapshot.universal_mid,
                    "pnl_at_breach": running_pnl,
                }
                realized_offset = running_pnl
                state_info = None
                waiting_reentry = True
                reentry_done = True

    if not rows:
        raise ValueError("No rows produced.")
    summary = {
        "final_pnl": rows[-1]["running_total_pnl"],
        "c2c_vol": scaled_volatility_or_none(c2c_var, session.config.market.calendar_days, session.config.market.intraday_var),
        "hedge_vol": scaled_volatility_or_none(hedge_var, session.config.market.calendar_days, session.config.market.intraday_var),
        "pnl_at_breach": breach["pnl_at_breach"] if breach else None,
        "breach_timestamp": breach["timestamp"].strftime("%Y-%m-%d %H:%M:%S") if breach else None,
        "breach_side": breach["side"] if breach else None,
        "breach_strike": breach["strike"] if breach else None,
        "max_net_delta_lots": max(net_deltas),
        "min_net_delta_lots": min(net_deltas),
        "avg_net_delta_lots": sum(net_deltas) / len(net_deltas),
        "synthetic_hedge_trades": hedge_trade_count,
        "synthetic_buy_lots": synthetic_buy_lots,
        "synthetic_sell_lots": synthetic_sell_lots,
        "synthetic_total_abs_lots": synthetic_buy_lots + synthetic_sell_lots,
        "premium_bought": premium_bought,
        "premium_sold": premium_sold,
        "initial_hedge_strike": initial_hedge_strike,
    }
    return {"rows": rows, "summary": summary}


def make_segment(
    session,
    price_book: DirectPriceBook,
    row: pd.Series,
    timestamp,
    user_value: int,
    hedge_strike_rule: str,
) -> SegmentInfo | None:
    snapshot = price_book.market_snapshot(row, timestamp, user_value)
    if snapshot is None:
        return None
    strikes = complete_strikes(price_book, row)
    portfolio = build_light_sample_portfolio(session, snapshot, strikes)
    if portfolio is None:
        return None
    frozen_ivs = capture_light_frozen_ivs(portfolio.positions, price_book, row, snapshot)
    if len(frozen_ivs) != len(portfolio.positions):
        return None
    hedge_strike = select_synthetic_strike(price_book, row, strikes, snapshot, hedge_strike_rule)
    short_calls = [p.strike for p in portfolio.positions if p.option_type == "CE" and float(p.lots) < 0]
    short_puts = [p.strike for p in portfolio.positions if p.option_type == "PE" and float(p.lots) < 0]
    return SegmentInfo(
        state=CheapestSyntheticHedgeState(
            portfolio=portfolio,
            options_pv_snapshot=None,
            hedge_strike=hedge_strike,
            candidate_strikes=strikes,
            hedge_strike_rule=hedge_strike_rule,
        ),
        short_call=max(short_calls) if short_calls else None,
        short_put=min(short_puts) if short_puts else None,
        hedge_strike=hedge_strike,
        user_value=snapshot.user_value,
        universal_mid=snapshot.universal_mid,
    )


def select_synthetic_strike(
    price_book: DirectPriceBook,
    row: pd.Series,
    strikes: list[int],
    snapshot,
    hedge_strike_rule: str,
) -> int:
    if hedge_strike_rule in {"cheapest", "cheapest_delayed"}:
        return cheapest_synthetic_strike(price_book, row, strikes, snapshot)
    if hedge_strike_rule == "closest_um":
        return closest_um_synthetic_strike(price_book, row, strikes, snapshot.universal_mid)
    if hedge_strike_rule in {"regular_delayed", "best_market", "best_market_delayed"}:
        return closest_um_synthetic_strike(price_book, row, strikes, snapshot.universal_mid)
    raise ValueError(f"Unsupported hedge strike rule: {hedge_strike_rule}")


def cheapest_synthetic_strike(price_book: DirectPriceBook, row: pd.Series, strikes: list[int], snapshot) -> int:
    candidates = []
    for strike in strikes:
        if abs(strike - snapshot.user_value) > price_book.market.synthetic_search_width:
            continue
        ce = price_book.option_price(row, strike, "CE")
        pe = price_book.option_price(row, strike, "PE")
        if ce is None or pe is None:
            continue
        candidates.append((ce + pe, abs(ce - pe), strike))
    if not candidates:
        raise ValueError("No complete synthetic strikes found.")
    return min(candidates)[2]


def closest_um_synthetic_strike(
    price_book: DirectPriceBook,
    row: pd.Series,
    strikes: list[int],
    universal_mid: float,
) -> int:
    candidates = []
    for strike in strikes:
        ce = price_book.option_price(row, strike, "CE")
        pe = price_book.option_price(row, strike, "PE")
        if ce is None or pe is None:
            continue
        candidates.append((abs(strike - universal_mid), strike))
    if not candidates:
        raise ValueError("No complete synthetic strikes found.")
    return min(candidates)[1]


def best_market_synthetic_strike(
    price_book: DirectPriceBook,
    row: pd.Series,
    snapshot,
    strikes: list[int],
    lots_change: float,
) -> int:
    candidates = []
    for strike in strikes:
        if abs(strike - snapshot.user_value) > price_book.market.synthetic_search_width:
            continue
        prices = price_book.hedge_prices(row, snapshot, strike)
        if prices is None:
            continue
        synth_bid, synth_ask, _ = prices
        if lots_change > 0:
            candidates.append((synth_ask, abs(strike - snapshot.universal_mid), strike))
        else:
            candidates.append((-synth_bid, abs(strike - snapshot.universal_mid), strike))
    if not candidates:
        raise ValueError("No eligible synthetic market strikes found.")
    return min(candidates)[2]


class CheapestSyntheticHedgeState(DirectHedgeState):
    def __init__(
        self,
        *args,
        candidate_strikes: list[int],
        hedge_strike_rule: str,
        **kwargs,
    ) -> None:
        super().__init__(*args, **kwargs)
        self.candidate_strikes = candidate_strikes
        self.hedge_strike_rule = hedge_strike_rule
        self.pending_hedge_trades = []

    @property
    def is_delayed_rule(self) -> bool:
        return self.hedge_strike_rule.endswith("_delayed")

    def update(
        self,
        price_book: DirectPriceBook,
        row: pd.Series,
        snapshot,
        timestamp,
        threshold_ratio: float,
    ) -> float | None:
        if self.is_delayed_rule:
            self.execute_pending_hedge_trade(price_book, row, snapshot, timestamp)
        values = direct_portfolio_values(
            self.portfolio.positions,
            price_book,
            row,
            snapshot,
            self.frozen_ivs,
        )
        if values is None:
            return None
        options_pv, delta_lots, gamma_lots = values
        self.last_options_pv = options_pv
        self.last_delta_lots = delta_lots
        self.last_gamma_lots = gamma_lots
        if self.options_pv_snapshot is None:
            self.options_pv_snapshot = options_pv
        threshold = dynamic_gamma_threshold(gamma_lots, threshold_ratio)
        combined_delta = delta_lots + self.cumulative_hedge_lots
        if abs(combined_delta) > threshold:
            self.add_hedge_trade(
                round(-combined_delta),
                price_book,
                row,
                snapshot,
                timestamp,
            )
        hedge_pnl = self.hedge_pnl(price_book, row, snapshot)
        self.last_total_pnl = (options_pv - self.options_pv_snapshot) + (hedge_pnl or 0.0)
        return self.last_total_pnl

    def execute_pending_hedge_trade(
        self,
        price_book: DirectPriceBook,
        row: pd.Series,
        snapshot,
        timestamp,
    ) -> None:
        if not self.pending_hedge_trades:
            return
        remaining = []
        for pending in self.pending_hedge_trades:
            hedge_strike = int(pending["hedge_strike"])
            lots_change = float(pending["lots_change"])
            prices = price_book.hedge_prices(row, snapshot, hedge_strike)
            if prices is None:
                remaining.append(pending)
                continue
            synth_bid, synth_ask, _ = prices
            trade_price = synth_ask if lots_change > 0 else synth_bid
            self.hedge_trades.append(
                {
                    "timestamp": timestamp,
                    "decision_timestamp": pending["decision_timestamp"],
                    "lots_change": lots_change,
                    "trade_price": trade_price,
                    "hedge_strike": hedge_strike,
                }
            )
            self.hedge_strike = hedge_strike
        self.pending_hedge_trades = remaining

    def add_hedge_trade(
        self,
        lots_change: float,
        price_book: DirectPriceBook,
        row: pd.Series,
        snapshot,
        timestamp,
    ) -> None:
        if abs(lots_change) < 1e-9:
            return
        if self.hedge_strike_rule in {"best_market", "best_market_delayed"}:
            hedge_strike = best_market_synthetic_strike(
                price_book,
                row,
                snapshot,
                self.candidate_strikes,
                lots_change,
            )
        elif self.hedge_strike_rule == "regular_delayed":
            hedge_strike = self.hedge_strike
        else:
            hedge_strike = select_synthetic_strike(
                price_book,
                row,
                self.candidate_strikes,
                snapshot,
                self.hedge_strike_rule,
            )
        if self.is_delayed_rule:
            self.pending_hedge_trades.append(
                {
                    "decision_timestamp": timestamp,
                    "lots_change": lots_change,
                    "hedge_strike": hedge_strike,
                }
            )
            self.cumulative_hedge_lots += lots_change
            self.hedge_strike = hedge_strike
            return
        prices = price_book.hedge_prices(row, snapshot, hedge_strike)
        if prices is None:
            return
        synth_bid, synth_ask, _ = prices
        trade_price = synth_ask if lots_change > 0 else synth_bid
        self.hedge_trades.append(
            {
                "timestamp": timestamp,
                "lots_change": lots_change,
                "trade_price": trade_price,
                "hedge_strike": hedge_strike,
            }
        )
        self.cumulative_hedge_lots += lots_change
        self.hedge_strike = hedge_strike

    def hedge_pnl(self, price_book: DirectPriceBook, row: pd.Series, snapshot) -> float | None:
        if not self.hedge_trades:
            return None
        multiplier = hedge_multiplier(self.portfolio.positions)
        total = 0.0
        for trade in self.hedge_trades:
            hedge_strike = int(trade.get("hedge_strike", self.hedge_strike))
            prices = price_book.hedge_prices(row, snapshot, hedge_strike)
            if prices is None:
                return None
            _, _, synth_mid = prices
            total += (
                float(trade["lots_change"])
                * (synth_mid - float(trade["trade_price"]))
                * multiplier
                / 1000
            )
        return total


def write_workbook(
    output: Path,
    underlying: str,
    trade_date,
    start_time: str,
    end_time: str,
    threshold_ratio: float,
    hedge_strike_rule: str,
    result,
    load_seconds: float,
    sim_seconds: float,
    total_seconds: float,
) -> None:
    summary = result["summary"]
    gammas = [row["gamma_lots"] for row in result["rows"] if row.get("gamma_lots") is not None]
    avg_gamma = sum(gammas) / len(gammas) if gammas else None
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    title_rules = {
        "cheapest": "Cheapest-Synthetic",
        "cheapest_delayed": "Cheapest Delayed Synthetic",
        "closest_um": "Closest-UM Synthetic",
        "regular_delayed": "Regular Delayed Synthetic",
        "best_market": "Best-Market Synthetic",
        "best_market_delayed": "Best-Market Delayed Synthetic",
    }
    rule_descriptions = {
        "cheapest": "Lowest CE+PE premium at each hedge event",
        "cheapest_delayed": (
            "Choose lowest CE+PE premium at T=t; apply hedge risk immediately and fill at T=t+1"
        ),
        "closest_um": "Complete strike closest to current UM at each hedge event",
        "regular_delayed": (
            "Use fixed segment hedge strike; apply hedge risk immediately and fill at T=t+1"
        ),
        "best_market": "Buy lowest synthetic ask; sell highest synthetic bid at each hedge event",
        "best_market_delayed": (
            "Decide best-market strike and lots at T=t; execute the stored trade at T=t+1"
        ),
    }
    title_rule = title_rules[hedge_strike_rule]
    rule_description = rule_descriptions[hedge_strike_rule]
    ws["A1"] = f"{underlying} {trade_date:%d-%b-%y} {title_rule} Restrike Diagnostics"
    ws["A1"].font = Font(bold=True, size=14)
    rows = [
        ("Underlying", underlying),
        ("Trade Date", trade_date.strftime("%d-%b-%y")),
        ("Reference Price", "Universal Mid"),
        ("Hedge Strike Rule", rule_description),
        ("Initial Hedge Strike", summary.get("initial_hedge_strike")),
        ("Start Time", start_time),
        ("End Time", end_time),
        ("Threshold Ratio", threshold_ratio),
        ("Restrike", "Yes"),
        ("Final PnL", summary.get("final_pnl")),
        ("c2c_vol", summary.get("c2c_vol")),
        ("hedge_vol", summary.get("hedge_vol")),
        ("PnL at Breach", summary.get("pnl_at_breach")),
        ("Breach Timestamp", summary.get("breach_timestamp")),
        ("Breach Side", summary.get("breach_side")),
        ("Breach Strike", summary.get("breach_strike")),
        ("Max Net Delta Lots", summary.get("max_net_delta_lots")),
        ("Min Net Delta Lots", summary.get("min_net_delta_lots")),
        ("Avg Net Delta Lots", summary.get("avg_net_delta_lots")),
        ("Average Gamma Lots", avg_gamma),
        ("Number of Hedge Events", summary.get("synthetic_hedge_trades")),
        ("Synthetic Buy Lots Hedged", summary.get("synthetic_buy_lots")),
        ("Synthetic Sell Lots Hedged", summary.get("synthetic_sell_lots")),
        ("Total Synthetic Lots Hedged", summary.get("synthetic_total_abs_lots")),
        ("Total Premium Bought", summary.get("premium_bought")),
        ("Total Premium Sold", summary.get("premium_sold")),
        ("Brokerage", brokerage(summary.get("premium_bought"), summary.get("premium_sold"))),
        ("Slippage", slippage(underlying, summary.get("synthetic_total_abs_lots"))),
        ("Data Load Seconds", load_seconds),
        ("Simulation Seconds", sim_seconds),
        ("Total Seconds", total_seconds),
    ]
    for row_idx, (label, value) in enumerate(rows, start=3):
        ws.cell(row_idx, 1).value = label
        ws.cell(row_idx, 1).font = Font(bold=True)
        ws.cell(row_idx, 2).value = value
        if isinstance(value, (int, float)):
            ws.cell(row_idx, 2).number_format = "#,##0" if label in {"Brokerage", "Slippage"} else "#,##0.00"

    path_ws = wb.create_sheet("Restrike Path")
    headers = [
        "timestamp",
        "universal_mid",
        "running_total_pnl",
        "net_delta_lots",
        "gamma_lots",
        "threshold_lots",
        "hedge_delta_lots",
        "traded_delta_lots",
        "hedge_strike",
    ]
    path_ws.append(headers)
    for cell in path_ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in result["rows"]:
        path_ws.append(
            [
                row["timestamp"].strftime("%Y-%m-%d %H:%M:%S"),
                row["universal_mid"],
                row["running_total_pnl"],
                row["net_delta_lots"],
                row["gamma_lots"],
                row["threshold_lots"],
                row["hedge_delta_lots"],
                row["traded_delta_lots"],
                row["hedge_strike"],
            ]
        )
    for data_row in path_ws.iter_rows(min_row=2, min_col=2, max_col=9):
        for cell in data_row:
            cell.number_format = "#,##0.00"
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A3" if sheet.title == "Summary" else "A2"
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 22
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def brokerage(premium_bought, premium_sold):
    if premium_bought is None or premium_sold is None:
        return None
    return 0.0024 * ((float(premium_bought) + float(premium_sold)) / 2)


def slippage(underlying: str, synthetic_lots):
    if synthetic_lots is None:
        return None
    normalized = underlying.strip().upper()
    if normalized == "NIFTY":
        return float(synthetic_lots) * 0.125 * 65
    if normalized == "SENSEX":
        return float(synthetic_lots) * 0.41 * 20
    return None


if __name__ == "__main__":
    main()
