from __future__ import annotations

from collections import defaultdict
from datetime import datetime, time, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(r"C:\Users\rishi\my_project\Backtest")
RUNS = ROOT / "runs_1s"
OUTPUT = RUNS / "pnl_interval_matrix_20260531.xlsx"

DATE_TEXTS = [
    "27-May-26",
    "26-May-26",
    "21-May-26",
    "19-May-26",
    "14-May-26",
    "12-May-26",
    "07-May-26",
    "05-May-26",
    "30-Apr-26",
    "28-Apr-26",
    "23-Apr-26",
    "21-Apr-26",
    "16-Apr-26",
    "13-Apr-26",
    "09-Apr-26",
    "07-Apr-26",
    "02-Apr-26",
    "30-Mar-26",
    "25-Mar-26",
    "24-Mar-26",
    "19-Mar-26",
    "17-Mar-26",
    "12-Mar-26",
    "10-Mar-26",
    "05-Mar-26",
    "02-Mar-26",
    "26-Feb-26",
    "24-Feb-26",
    "19-Feb-26",
    "17-Feb-26",
    "12-Feb-26",
    "10-Feb-26",
    "05-Feb-26",
    "03-Feb-26",
    "29-Jan-26",
    "27-Jan-26",
    "22-Jan-26",
    "20-Jan-26",
    "14-Jan-26",
    "13-Jan-26",
    "08-Jan-26",
    "06-Jan-26",
    "01-Jan-26",
    "30-Dec-25",
    "24-Dec-25",
    "23-Dec-25",
    "18-Dec-25",
    "16-Dec-25",
    "11-Dec-25",
    "09-Dec-25",
    "04-Dec-25",
    "02-Dec-25",
    "27-Nov-25",
    "25-Nov-25",
    "20-Nov-25",
    "18-Nov-25",
    "13-Nov-25",
    "11-Nov-25",
    "06-Nov-25",
    "04-Nov-25",
    "30-Oct-25",
    "28-Oct-25",
    "23-Oct-25",
    "20-Oct-25",
    "16-Oct-25",
    "14-Oct-25",
    "09-Oct-25",
    "07-Oct-25",
    "01-Oct-25",
]


def main() -> None:
    latest = find_latest_pnl_workbooks()
    intervals = interval_labels()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    coverage_rows = []
    for underlying in ("NIFTY", "SENSEX"):
        ws = wb.create_sheet(f"{underlying} Frozen IV")
        coverage_rows.extend(write_matrix(ws, underlying, latest, intervals))
    coverage = wb.create_sheet("Coverage")
    write_coverage(coverage, coverage_rows)
    wb.save(OUTPUT)
    print(OUTPUT)


def find_latest_pnl_workbooks() -> dict[tuple[str, str], Path]:
    candidates: dict[tuple[str, str], list[Path]] = defaultdict(list)
    for path in RUNS.rglob("*_pnl_only.xlsx"):
        name = path.name.lower()
        parts = name.split("_")
        if len(parts) < 4 or parts[2] != "20"[:0]:
            pass
        if name.startswith("nifty_1s_") or name.startswith("sensex_1s_"):
            underlying = name.split("_")[0].upper()
            date_key = name.split("_")[2]
            if len(date_key) == 8 and date_key.isdigit():
                candidates[(underlying, date_key)].append(path)
    return {
        key: sorted(paths, key=lambda item: item.stat().st_mtime)[-1]
        for key, paths in candidates.items()
    }


def write_matrix(ws, underlying: str, workbooks: dict[tuple[str, str], Path], intervals):
    ws["A1"] = f"{underlying} Frozen IV Running PnL: 10-minute interval changes"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(intervals) + 2)
    headers = ["Date", *intervals, "Final Frozen IV PnL"]
    ws.append([])
    ws.append(headers)
    style_header(ws, 3, len(headers))

    coverage_rows = []
    matrix_values = []
    for text in DATE_TEXTS:
        date_key = datetime.strptime(text, "%d-%b-%y").strftime("%Y%m%d")
        path = workbooks.get((underlying, date_key))
        if path is None:
            values = [None] * len(intervals)
            final_value = None
            status = "missing time series"
        else:
            values, final_value = interval_changes(path, intervals)
            status = "loaded"
        ws.append([text, *values, final_value])
        matrix_values.append(values)
        coverage_rows.append([underlying, text, status, str(path) if path else ""])

    avg_row = ws.max_row + 2
    ws.cell(avg_row, 1).value = "Average"
    ws.cell(avg_row, 1).font = Font(bold=True)
    for col in range(2, len(intervals) + 2):
        values = [
            row[col - 2]
            for row in matrix_values
            if isinstance(row[col - 2], (int, float))
        ]
        ws.cell(avg_row, col).value = sum(values) / len(values) if values else None
    final_values = [
        ws.cell(row, len(headers)).value
        for row in range(4, 4 + len(DATE_TEXTS))
        if isinstance(ws.cell(row, len(headers)).value, (int, float))
    ]
    ws.cell(avg_row, len(headers)).value = (
        sum(final_values) / len(final_values) if final_values else None
    )
    ws.cell(avg_row, len(headers)).font = Font(bold=True)

    format_matrix(ws, len(headers), avg_row)
    return coverage_rows


def interval_changes(path: Path, intervals: list[str]) -> tuple[list[float | None], float | None]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["PnL Timeseries"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    points = []
    for timestamp, _um, _pnl, frozen_pnl in rows:
        if timestamp is None or frozen_pnl is None:
            continue
        dt = datetime.strptime(str(timestamp), "%Y-%m-%d %H:%M:%S")
        points.append((dt, float(frozen_pnl)))
    if not points:
        return [None] * len(intervals), None
    by_time = {dt.time(): pnl for dt, pnl in points}
    values = []
    for label in intervals:
        start_text, end_text = label.split("-")
        start = parse_clock(start_text)
        end = parse_clock(end_text)
        start_value = by_time.get(start)
        end_value = by_time.get(end)
        if start_value is None or end_value is None:
            values.append(None)
        else:
            values.append(round(end_value - start_value, 2))
    return values, round(points[-1][1], 2)


def interval_labels() -> list[str]:
    labels = []
    current = datetime.combine(datetime.today(), time(9, 20))
    final = datetime.combine(datetime.today(), time(15, 20))
    while current < final:
        nxt = current + timedelta(minutes=10)
        labels.append(f"{current:%H:%M}-{nxt:%H:%M}")
        current = nxt
    return labels


def parse_clock(value: str) -> time:
    hour, minute = value.split(":")
    return time(int(hour), int(minute), 0)


def style_header(ws, row: int, columns: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for col in range(1, columns + 1):
        cell = ws.cell(row, col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")


def format_matrix(ws, columns: int, avg_row: int) -> None:
    ws.freeze_panes = "B4"
    ws.column_dimensions["A"].width = 13
    for col in range(2, columns + 1):
        ws.column_dimensions[get_column_letter(col)].width = 12
    for row in ws.iter_rows(min_row=4, max_row=avg_row, min_col=2, max_col=columns):
        for cell in row:
            cell.number_format = "#,##0.00"
    fill = PatternFill("solid", fgColor="D9EAF7")
    for col in range(1, columns + 1):
        cell = ws.cell(avg_row, col)
        cell.fill = fill
        cell.font = Font(bold=True)


def write_coverage(ws, rows) -> None:
    ws.append(["Underlying", "Date", "Status", "Workbook"])
    style_header(ws, 1, 4)
    for row in rows:
        ws.append(row)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 13
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 110


if __name__ == "__main__":
    main()
