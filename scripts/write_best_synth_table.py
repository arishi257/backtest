from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT_DIR = ROOT / "best_synth_sims"

VERSION_FILES = {
    "Regular": "{prefix}_restrike_diagnostics.xlsx",
    "Regular Delayed": "{prefix}_regular_delayed_synth_restrike_diagnostics.xlsx",
    "Cheapest Synth": "{prefix}_cheapest_synth_restrike_diagnostics.xlsx",
    "Cheapest Delayed": "{prefix}_cheapest_delayed_synth_restrike_diagnostics.xlsx",
    "Best Market Synth": "{prefix}_best_market_synth_restrike_diagnostics.xlsx",
    "Best Market Delayed": "{prefix}_best_market_delayed_synth_restrike_diagnostics.xlsx",
}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--date", required=True, help="Date label like 06Jan26.")
    parser.add_argument("--underlying", required=True, choices=("NIFTY", "SENSEX"))
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    args = parser.parse_args()

    prefix = f"{args.underlying.lower()}_{args.date}_1000_threshold_0p1"
    files = {
        version: args.output_dir / template.format(prefix=prefix)
        for version, template in VERSION_FILES.items()
    }
    existing = {version: path for version, path in files.items() if path.exists()}
    if not existing:
        raise FileNotFoundError(f"No comparison workbooks found for {prefix} in {args.output_dir}.")

    headers = [
        "Version",
        "Final PnL",
        "Hedge Events",
        "Synth Lots",
        "Premium Bought",
        "Premium Sold",
        "Brokerage",
        "Slippage",
    ]
    rows = []
    for version, path in existing.items():
        summary = read_summary(path)
        rows.append(
            [
                version,
                summary.get("Final PnL"),
                summary.get("Number of Hedge Events"),
                summary.get("Total Synthetic Lots Hedged"),
                summary.get("Total Premium Bought"),
                summary.get("Total Premium Sold"),
                summary.get("Brokerage"),
                summary.get("Slippage"),
            ]
        )

    title = f"{args.underlying} {format_date_label(args.date)} 10am Restrike Comparison"
    for path in existing.values():
        write_table_sheet(path, title, headers, rows)

    for row in rows:
        formatted = [str(row[0])] + [
            f"{value:,.0f}" if isinstance(value, (int, float)) else str(value)
            for value in row[1:]
        ]
        print("| " + " | ".join(formatted) + " |")


def read_summary(path: Path) -> dict[str, object]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook["Summary"]
    data = {}
    for row in worksheet.iter_rows(values_only=True):
        if row and row[0] is not None:
            data[str(row[0])] = row[1]
    workbook.close()
    return data


def write_table_sheet(path: Path, title: str, headers: list[str], rows: list[list[object]]) -> None:
    workbook = load_workbook(path)
    if "table" in workbook.sheetnames:
        del workbook["table"]
    worksheet = workbook.create_sheet("table", 0)
    worksheet["A1"] = title
    worksheet["A1"].font = Font(bold=True, size=14)
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    worksheet.append([])
    worksheet.append(headers)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    subtle_fill = PatternFill("solid", fgColor="D9EAF7")
    white_font = Font(color="FFFFFF", bold=True)
    bold = Font(bold=True)
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in worksheet[3]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    for row in rows:
        worksheet.append(row)

    for row in worksheet.iter_rows(min_row=4, max_row=3 + len(rows), min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="right" if cell.column > 1 else "left")
        row[0].font = bold

    for col in range(2, len(headers) + 1):
        for col_cells in worksheet.iter_cols(min_col=col, max_col=col, min_row=4, max_row=3 + len(rows)):
            for cell in col_cells:
                cell.number_format = "#,##0"

    for cell in worksheet[4]:
        cell.fill = subtle_fill

    worksheet.freeze_panes = "A4"
    widths = [24, 16, 14, 14, 18, 18, 14, 14]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    workbook.save(path)
    workbook.close()


def format_date_label(value: str) -> str:
    return f"{value[:2]}-{value[2:5]}-{value[5:]}"


if __name__ == "__main__":
    main()
