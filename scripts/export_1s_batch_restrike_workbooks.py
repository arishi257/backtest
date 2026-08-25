from __future__ import annotations

import argparse
import math
import sys
import time
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
SCRIPTS = ROOT / "scripts"
for path in (SRC, SCRIPTS):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from backtest.config import DEFAULT_WORKBOOK, normalize_underlying
from backtest.replay import CsvReplayFeed, register_token_tickers
from backtest.sessions import build_backtest_sessions
from backtest_sample_hf.__main__ import load_sample_hf_option_dataset, parse_date_key
from export_sample_hf_diagnostics import close_to_close_variance, scaled_volatility_or_none
from run_nifty_1m_light_batch import complete_strikes
from run_sample_hf_pnl_only import (
    DirectHedgeState,
    DirectPriceBook,
    build_light_sample_portfolio,
    capture_light_frozen_ivs,
    hedge_multiplier,
)


DEFAULT_DATA_DIR = Path(r"C:\options data\1s data")
DEFAULT_OUTPUT_DIR = Path(r"C:\Users\rishi\OneDrive\Desktop\Summary")
START_TIME = "09:20"
END_TIME = "15:15"
THRESHOLD_RATIO = 0.40


@dataclass
class SegmentInfo:
    state: DirectHedgeState
    short_call: int | None
    short_put: int | None
    hedge_strike: int
    user_value: int
    universal_mid: float


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-dir", type=Path, default=DEFAULT_DATA_DIR)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--threshold-ratio", type=float, default=THRESHOLD_RATIO)
    parser.add_argument("--start-time", default=START_TIME)
    parser.add_argument("--end-time", default=END_TIME)
    args = parser.parse_args()

    jobs = [("NIFTY", date(2026, 5, 26)), ("SENSEX", date(2026, 5, 27))]
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for underlying, trade_date in jobs:
        started = time.perf_counter()
        cache_start = time.perf_counter()
        dataset = load_sample_hf_option_dataset(args.data_dir, trade_date, underlying)
        load_seconds = time.perf_counter() - cache_start

        no_restrike = run_path(
            dataset,
            DEFAULT_WORKBOOK,
            args.threshold_ratio,
            args.start_time,
            args.end_time,
            with_restrike=False,
        )
        restrike = run_path(
            dataset,
            DEFAULT_WORKBOOK,
            args.threshold_ratio,
            args.start_time,
            args.end_time,
            with_restrike=True,
        )
        output = args.output_dir / f"{underlying.lower()}_{trade_date:%d%b%y}_1s_light_restrike.xlsx"
        write_workbook(output, underlying, trade_date, no_restrike, restrike, load_seconds)
        print(
            f"{underlying} {trade_date:%d-%b-%y}: "
            f"no_restrike={no_restrike['summary']['final_pnl']:.0f}, "
            f"restrike={restrike['summary']['final_pnl']:.0f}, "
            f"workbook={output}, seconds={time.perf_counter() - started:.2f}"
        )


def run_path(dataset, workbook: Path, threshold_ratio: float, start_time: str, end_time: str, with_restrike: bool):
    replay = None
    current_time = lambda: replay.now()
    sessions = build_backtest_sessions(dataset, workbook, current_time, refresh_ms=0)
    replay = CsvReplayFeed(dataset, sessions)
    register_token_tickers(sessions, dataset)
    if len(sessions) != 1:
        raise ValueError("Expected one expiry session.")
    session = sessions[0]
    surface = dataset.option_surface(session.spec.expiry)
    price_book = DirectPriceBook(dataset, surface, session)
    multiplier = None

    state_info: SegmentInfo | None = None
    waiting_reentry = False
    reentry_done = False
    user_value = session.config.market.user_value
    realized_offset = 0.0
    breach = None
    rows = []
    prev_um = None
    prev_hedge_um = None
    c2c_var = 0.0
    hedge_var = 0.0
    net_deltas = []
    hedge_trade_count = 0
    synthetic_buy_lots = 0.0
    synthetic_sell_lots = 0.0
    premium_bought = 0.0
    premium_sold = 0.0

    while replay.advance():
        timestamp = replay.now()
        if not in_time_window(timestamp, start_time, end_time):
            continue
        row = surface.loc[dataset.timestamps[replay.position]]

        if waiting_reentry:
            state_info = make_segment(session, price_book, row, timestamp, user_value)
            if state_info is None:
                continue
            multiplier = hedge_multiplier(state_info.state.portfolio.positions)
            bought, sold = portfolio_premiums(state_info.state.portfolio.positions, price_book, row, multiplier)
            premium_bought += bought
            premium_sold += sold
            user_value = state_info.user_value
            waiting_reentry = False

        if state_info is None:
            state_info = make_segment(session, price_book, row, timestamp, user_value)
            if state_info is None:
                continue
            multiplier = hedge_multiplier(state_info.state.portfolio.positions)
            bought, sold = portfolio_premiums(state_info.state.portfolio.positions, price_book, row, multiplier)
            premium_bought += bought
            premium_sold += sold
            user_value = state_info.user_value

        snapshot = price_book.market_snapshot(row, timestamp, user_value)
        if snapshot is None:
            continue
        user_value = snapshot.user_value
        prior_trade_count = len(state_info.state.hedge_trades)
        segment_pnl = state_info.state.update(price_book, row, snapshot, timestamp, threshold_ratio)
        if segment_pnl is None:
            continue
        new_trades = state_info.state.hedge_trades[prior_trade_count:]
        for trade in new_trades:
            lots_change = float(trade["lots_change"])
            if lots_change > 0:
                synthetic_buy_lots += lots_change
            elif lots_change < 0:
                synthetic_sell_lots += abs(lots_change)
            bought, sold = synthetic_trade_premiums(
                price_book,
                row,
                state_info.hedge_strike,
                lots_change,
                multiplier,
            )
            premium_bought += bought
            premium_sold += sold
        hedge_trade_count += len(new_trades)

        running_pnl = realized_offset + segment_pnl
        net_delta = (state_info.state.last_delta_lots or 0.0) + state_info.state.cumulative_hedge_lots
        gamma_lots = state_info.state.last_gamma_lots
        net_deltas.append(net_delta)
        traded_um = snapshot.universal_mid if new_trades else None
        if prev_um is not None:
            c2c_var += close_to_close_variance(prev_um, snapshot.universal_mid)
        prev_um = snapshot.universal_mid
        if traded_um is not None:
            if prev_hedge_um is not None:
                hedge_var += close_to_close_variance(prev_hedge_um, traded_um)
            prev_hedge_um = traded_um

        rows.append(
            {
                "timestamp": timestamp,
                "universal_mid": snapshot.universal_mid,
                "running_total_pnl": running_pnl,
                "net_delta_lots": net_delta,
                "gamma_lots": gamma_lots,
            }
        )

        if with_restrike and not reentry_done:
            breach_side = None
            breach_strike = None
            if state_info.short_call is not None and snapshot.universal_mid > state_info.short_call:
                breach_side = "CALL"
                breach_strike = state_info.short_call
            elif state_info.short_put is not None and snapshot.universal_mid < state_info.short_put:
                breach_side = "PUT"
                breach_strike = state_info.short_put
            if breach_side:
                bought, sold = portfolio_liquidation_premiums(
                    state_info.state.portfolio.positions,
                    price_book,
                    row,
                    multiplier,
                )
                premium_bought += bought
                premium_sold += sold
                breach = {
                    "side": breach_side,
                    "timestamp": timestamp,
                    "strike": breach_strike,
                    "universal_mid": snapshot.universal_mid,
                    "pnl_at_breach": running_pnl,
                }
                realized_offset = running_pnl
                state_info = None
                waiting_reentry = True
                reentry_done = True

    if not rows:
        raise ValueError("No rows produced.")
    summary = {
        "final_pnl": rows[-1]["running_total_pnl"],
        "c2c_vol": scaled_volatility_or_none(c2c_var, session.config.market.calendar_days, session.config.market.intraday_var),
        "hedge_vol": scaled_volatility_or_none(hedge_var, session.config.market.calendar_days, session.config.market.intraday_var),
        "pnl_at_breach": breach["pnl_at_breach"] if breach else None,
        "breach_timestamp": breach["timestamp"].strftime("%Y-%m-%d %H:%M:%S") if breach else None,
        "breach_side": breach["side"] if breach else None,
        "breach_strike": breach["strike"] if breach else None,
        "max_net_delta_lots": max(net_deltas),
        "min_net_delta_lots": min(net_deltas),
        "avg_net_delta_lots": sum(net_deltas) / len(net_deltas),
        "synthetic_hedge_trades": hedge_trade_count,
        "synthetic_buy_lots": synthetic_buy_lots,
        "synthetic_sell_lots": synthetic_sell_lots,
        "synthetic_total_abs_lots": synthetic_buy_lots + synthetic_sell_lots,
        "premium_bought": premium_bought,
        "premium_sold": premium_sold,
    }
    return {"rows": rows, "summary": summary}


def make_segment(session, price_book: DirectPriceBook, row: pd.Series, timestamp: datetime, user_value: int) -> SegmentInfo | None:
    snapshot = price_book.market_snapshot(row, timestamp, user_value)
    if snapshot is None:
        return None
    strikes = complete_strikes(price_book, row)
    portfolio = build_light_sample_portfolio(session, snapshot, strikes)
    if portfolio is None:
        return None
    frozen_ivs = capture_light_frozen_ivs(portfolio.positions, price_book, row, snapshot)
    if len(frozen_ivs) != len(portfolio.positions):
        return None
    hedge_strike = min(strikes, key=lambda strike: (abs(strike - snapshot.universal_mid), strike))
    short_calls = [position.strike for position in portfolio.positions if position.option_type == "CE" and float(position.lots) < 0]
    short_puts = [position.strike for position in portfolio.positions if position.option_type == "PE" and float(position.lots) < 0]
    return SegmentInfo(
        state=DirectHedgeState(portfolio=portfolio, options_pv_snapshot=None, hedge_strike=hedge_strike),
        short_call=max(short_calls) if short_calls else None,
        short_put=min(short_puts) if short_puts else None,
        hedge_strike=hedge_strike,
        user_value=snapshot.user_value,
        universal_mid=snapshot.universal_mid,
    )


def option_price(price_book: DirectPriceBook, row: pd.Series, strike: int, option_type: str) -> float:
    value = price_book.option_price(row, strike, option_type)
    if value is None:
        raise ValueError(f"Missing {strike} {option_type} price.")
    return value


def portfolio_premiums(positions, price_book: DirectPriceBook, row: pd.Series, multiplier: float) -> tuple[float, float]:
    bought = 0.0
    sold = 0.0
    for position in positions:
        premium = option_price(price_book, row, position.strike, position.option_type) * abs(float(position.lots)) * multiplier
        if float(position.lots) > 0:
            bought += premium
        else:
            sold += premium
    return bought, sold


def portfolio_liquidation_premiums(positions, price_book: DirectPriceBook, row: pd.Series, multiplier: float) -> tuple[float, float]:
    bought = 0.0
    sold = 0.0
    for position in positions:
        premium = option_price(price_book, row, position.strike, position.option_type) * abs(float(position.lots)) * multiplier
        if float(position.lots) < 0:
            bought += premium
        else:
            sold += premium
    return bought, sold


def synthetic_trade_premiums(price_book: DirectPriceBook, row: pd.Series, strike: int, lots_change: float, multiplier: float) -> tuple[float, float]:
    if abs(lots_change) < 1e-9:
        return 0.0, 0.0
    call = option_price(price_book, row, strike, "CE") * abs(lots_change) * multiplier
    put = option_price(price_book, row, strike, "PE") * abs(lots_change) * multiplier
    if lots_change > 0:
        return call, put
    return put, call


def in_time_window(timestamp: datetime, start_time: str, end_time: str) -> bool:
    hhmm = timestamp.strftime("%H:%M")
    if hhmm < start_time or hhmm > end_time:
        return False
    return not (hhmm == end_time and timestamp.second > 0)


def write_workbook(path: Path, underlying: str, trade_date: date, no_restrike, restrike, load_seconds: float) -> None:
    wb = openpyxl.Workbook()
    summary_ws = wb.active
    summary_ws.title = "Summary"
    write_summary(summary_ws, underlying, trade_date, no_restrike, restrike, load_seconds)
    write_path_sheet(wb.create_sheet("No Restrike Path"), no_restrike["rows"])
    write_path_sheet(wb.create_sheet("Restrike Path"), restrike["rows"])
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for column in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(column)].width = 18
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_summary(ws, underlying: str, trade_date: date, no_restrike, restrike, load_seconds: float) -> None:
    ws["A1"] = f"{underlying} {trade_date:%d-%b-%y} 1s Light Summary"
    ws["A1"].font = Font(bold=True, size=14)
    rows = [
        ("Data Load Seconds", load_seconds, None),
        ("Metric", "No Restrike", "Restrike"),
        ("Final PnL", no_restrike["summary"]["final_pnl"], restrike["summary"]["final_pnl"]),
        ("c2c_vol", no_restrike["summary"]["c2c_vol"], restrike["summary"]["c2c_vol"]),
        ("hedge_vol", no_restrike["summary"]["hedge_vol"], restrike["summary"]["hedge_vol"]),
        ("PnL at Breach", no_restrike["summary"]["pnl_at_breach"], restrike["summary"]["pnl_at_breach"]),
        ("Breach Timestamp", no_restrike["summary"]["breach_timestamp"], restrike["summary"]["breach_timestamp"]),
        ("Breach Side", no_restrike["summary"]["breach_side"], restrike["summary"]["breach_side"]),
        ("Breach Strike", no_restrike["summary"]["breach_strike"], restrike["summary"]["breach_strike"]),
        ("Max Net Delta Lots", no_restrike["summary"]["max_net_delta_lots"], restrike["summary"]["max_net_delta_lots"]),
        ("Min Net Delta Lots", no_restrike["summary"]["min_net_delta_lots"], restrike["summary"]["min_net_delta_lots"]),
        ("Avg Net Delta Lots", no_restrike["summary"]["avg_net_delta_lots"], restrike["summary"]["avg_net_delta_lots"]),
        ("Synthetic Hedge Trades", no_restrike["summary"]["synthetic_hedge_trades"], restrike["summary"]["synthetic_hedge_trades"]),
        ("Total Premium Bought", no_restrike["summary"]["premium_bought"], restrike["summary"]["premium_bought"]),
        ("Total Premium Sold", no_restrike["summary"]["premium_sold"], restrike["summary"]["premium_sold"]),
    ]
    for r, row in enumerate(rows, start=3):
        for c, value in enumerate(row, start=1):
            ws.cell(r, c).value = value
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[4]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for row in ws.iter_rows(min_row=5, max_col=3):
        if isinstance(row[1].value, (int, float)):
            row[1].number_format = "#,##0.00"
        if isinstance(row[2].value, (int, float)):
            row[2].number_format = "#,##0.00"


def write_path_sheet(ws, rows) -> None:
    headers = ["timestamp", "universal_mid", "running_total_pnl", "net_delta_lots", "gamma_lots"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in rows:
        ws.append(
            [
                row["timestamp"].strftime("%Y-%m-%d %H:%M:%S"),
                row["universal_mid"],
                row["running_total_pnl"],
                row["net_delta_lots"],
                row["gamma_lots"],
            ]
        )
    for data_row in ws.iter_rows(min_row=2, min_col=2, max_col=5):
        for cell in data_row:
            cell.number_format = "#,##0.00"


if __name__ == "__main__":
    main()
