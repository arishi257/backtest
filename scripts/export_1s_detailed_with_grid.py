from __future__ import annotations

import argparse
import sys
import time
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
SCRIPTS = ROOT / "scripts"
for path in (SRC, SCRIPTS):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from backtest.config import DEFAULT_WORKBOOK, normalize_underlying
from backtest_sample_hf.__main__ import load_sample_hf_option_dataset, parse_date_key
from export_1s_batch_restrike_workbooks import run_path, write_workbook


DEFAULT_DATA_DIR = Path(r"C:\options data\1s data")
DEFAULT_OUTPUT_DIR = Path(r"C:\Users\rishi\OneDrive\Desktop\Summary")
GRID_START_TIMES = ("09:20", "10:00")
GRID_THRESHOLDS = (0.1, 0.4, 0.9)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--date", required=True)
    parser.add_argument("--underlying", required=True, type=normalize_underlying, choices=("NIFTY", "SENSEX"))
    parser.add_argument("--data-dir", type=Path, default=DEFAULT_DATA_DIR)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--start-time", default="09:20")
    parser.add_argument("--end-time", default="15:15")
    parser.add_argument("--threshold-ratio", type=float, default=0.40)
    args = parser.parse_args()

    trade_date = parse_date_key(args.date)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    output = args.output_dir / f"{args.underlying.lower()}_{trade_date:%d%b%y}_1s_light_detailed_with_grid.xlsx"

    total_start = time.perf_counter()
    load_start = time.perf_counter()
    dataset = load_sample_hf_option_dataset(args.data_dir, trade_date, args.underlying)
    load_seconds = time.perf_counter() - load_start

    detailed_no_restrike = run_path(
        dataset,
        DEFAULT_WORKBOOK,
        args.threshold_ratio,
        args.start_time,
        args.end_time,
        with_restrike=False,
    )
    detailed_restrike = run_path(
        dataset,
        DEFAULT_WORKBOOK,
        args.threshold_ratio,
        args.start_time,
        args.end_time,
        with_restrike=True,
    )
    write_workbook(output, args.underlying, trade_date, detailed_no_restrike, detailed_restrike, load_seconds)

    grid_results = []
    for start_time in GRID_START_TIMES:
        for threshold in GRID_THRESHOLDS:
            for restrike in (False, True):
                if (
                    start_time == args.start_time
                    and abs(threshold - args.threshold_ratio) < 1e-12
                    and not restrike
                ):
                    result = detailed_no_restrike
                    seconds = None
                elif (
                    start_time == args.start_time
                    and abs(threshold - args.threshold_ratio) < 1e-12
                    and restrike
                ):
                    result = detailed_restrike
                    seconds = None
                else:
                    scenario_start = time.perf_counter()
                    result = run_path(
                        dataset,
                        DEFAULT_WORKBOOK,
                        threshold,
                        start_time,
                        args.end_time,
                        with_restrike=restrike,
                    )
                    seconds = time.perf_counter() - scenario_start
                grid_results.append(
                    {
                        "start_time": start_time,
                        "threshold": threshold,
                        "restrike": "Yes" if restrike else "No",
                        "final_pnl": result["summary"]["final_pnl"],
                        "seconds": seconds,
                    }
                )

    append_grid_sheet(output, args.underlying, trade_date, args.end_time, grid_results)
    total_seconds = time.perf_counter() - total_start

    print(f"OUTPUT,{output}")
    print(f"TOTAL_SECONDS,{total_seconds:.2f}")
    print("START,THRESHOLD,NO_RESTRIKE,RESTRIKE")
    for start_time in GRID_START_TIMES:
        for threshold in GRID_THRESHOLDS:
            no_restrike = next(
                row["final_pnl"]
                for row in grid_results
                if row["start_time"] == start_time and row["threshold"] == threshold and row["restrike"] == "No"
            )
            restrike = next(
                row["final_pnl"]
                for row in grid_results
                if row["start_time"] == start_time and row["threshold"] == threshold and row["restrike"] == "Yes"
            )
            print(f"{start_time},{threshold:.1f},{no_restrike:.2f},{restrike:.2f}")


def append_grid_sheet(output: Path, underlying: str, trade_date: date, end_time: str, results: list[dict[str, object]]) -> None:
    wb = openpyxl.load_workbook(output)
    if "Scenario Grid" in wb.sheetnames:
        del wb["Scenario Grid"]
    ws = wb.create_sheet("Scenario Grid")
    ws["A1"] = f"{underlying} {trade_date:%d-%b-%y} Scenario Grid"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "End Time"
    ws["B2"] = end_time

    headers = ["Start Time", "Threshold", "Restrike", "Final PnL", "Scenario Seconds"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(4, col)
        cell.value = header
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row_idx, row in enumerate(results, start=5):
        ws.cell(row_idx, 1).value = row["start_time"]
        ws.cell(row_idx, 2).value = row["threshold"]
        ws.cell(row_idx, 3).value = row["restrike"]
        ws.cell(row_idx, 4).value = row["final_pnl"]
        ws.cell(row_idx, 5).value = row["seconds"]
        ws.cell(row_idx, 4).number_format = "#,##0.00"
        ws.cell(row_idx, 5).number_format = "#,##0.00"

    matrix_row = 20
    ws.cell(matrix_row, 1).value = "Final PnL Matrix"
    ws.cell(matrix_row, 1).font = Font(bold=True, size=12)
    matrix_row += 2
    for start_time in GRID_START_TIMES:
        ws.cell(matrix_row, 1).value = f"Start {start_time}"
        ws.cell(matrix_row, 1).font = Font(bold=True)
        matrix_row += 1
        for col, header in enumerate(("Threshold", "No Restrike", "Restrike"), start=1):
            cell = ws.cell(matrix_row, col)
            cell.value = header
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        matrix_row += 1
        for threshold in GRID_THRESHOLDS:
            ws.cell(matrix_row, 1).value = threshold
            for col, restrike in ((2, "No"), (3, "Yes")):
                match = next(
                    row
                    for row in results
                    if row["start_time"] == start_time
                    and row["threshold"] == threshold
                    and row["restrike"] == restrike
                )
                ws.cell(matrix_row, col).value = match["final_pnl"]
                ws.cell(matrix_row, col).number_format = "#,##0.00"
            matrix_row += 1
        matrix_row += 2

    ws.freeze_panes = "A5"
    for sheet in wb.worksheets:
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 18
    wb.save(output)


if __name__ == "__main__":
    main()
