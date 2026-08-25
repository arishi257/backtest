from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_PATH = PROJECT_ROOT / "outputs" / "sensex_5s_20260527_portfolio_greeks.xlsx"
OUTPUT_PATH = PROJECT_ROOT / "outputs" / "sensex_5s_20260527_net_delta.xlsx"


def main() -> None:
    source = load_workbook(SOURCE_PATH, data_only=True)
    source_sheet = source["Time Slices"]
    headers = [source_sheet.cell(1, column).value for column in range(1, source_sheet.max_column + 1)]
    timestamp_col = headers.index("timestamp") + 1
    net_delta_col = headers.index("live_combined_delta_lots") + 1
    gamma_col = headers.index("live_gamma_lots_10bps") + 1

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Net Delta"
    sheet.append(["timestamp", "net_delta_lots", "gamma_lots"])
    for row_index in range(2, source_sheet.max_row + 1):
        sheet.append(
            [
                source_sheet.cell(row_index, timestamp_col).value,
                source_sheet.cell(row_index, net_delta_col).value,
                source_sheet.cell(row_index, gamma_col).value,
            ]
        )

    format_sheet(sheet)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT_PATH)
    print(OUTPUT_PATH)
    print(sheet.max_row - 1)


def format_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    if sheet.max_row > 1:
        table = Table(displayName="Sensex5sNetDelta", ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)

    for row_index in range(2, sheet.max_row + 1):
        sheet[f"A{row_index}"].number_format = "yyyy-mm-dd hh:mm:ss"
        sheet[f"B{row_index}"].number_format = "#,##0.00"
        sheet[f"C{row_index}"].number_format = "#,##0.00"

    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
        sheet.column_dimensions[letter].width = min(max(max_length + 2, 14), 24)


if __name__ == "__main__":
    main()
