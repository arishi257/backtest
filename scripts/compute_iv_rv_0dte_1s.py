from __future__ import annotations

import argparse
import csv
import math
import re
import sys
import time
from datetime import date, datetime
from pathlib import Path

import pandas as pd
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
SCRIPTS = ROOT / "scripts"
for path in (SRC, SCRIPTS):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from backtest.config import DEFAULT_WORKBOOK, normalize_underlying
from backtest.replay import CsvReplayFeed, register_token_tickers
from backtest.sessions import build_backtest_sessions
from export_sample_hf_diagnostics import close_to_close_variance, scaled_volatility_or_none
from fit_sensex.pricing.black_scholes import implied_volatility
from light_run import load_breeze_parquet_option_dataset
from run_nifty_1m_light_batch import complete_strikes
from run_sample_hf_pnl_only import DirectPriceBook, build_light_sample_portfolio, direct_portfolio_values
from backtest_sample_hf.__main__ import load_sample_hf_option_dataset


SAMPLE_ROOT = Path(r"C:\options data\1s data")
BREEZE_ROOT = Path(r"C:\Users\rishi\my_project\historical\breeze 1s")
OUTPUT_DIR = ROOT / "runs_1s" / "analysis"
OPTION_FILE_RE = re.compile(
    r"^(?P<underlying>NIFTY|SENSEX)_(?P<expiry>\d{2}[A-Z]{3}\d{2})_"
    r"(?P<strike>\d+)_(?P<option_type>CE|PE)\.csv$",
    re.IGNORECASE,
)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--sample-root", type=Path, default=SAMPLE_ROOT)
    parser.add_argument("--breeze-root", type=Path, default=BREEZE_ROOT)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    parser.add_argument("--start-time", default="09:20")
    parser.add_argument("--end-time", default="15:15")
    parser.add_argument("--preload-time", default="09:15")
    args = parser.parse_args()

    args.output_dir.mkdir(parents=True, exist_ok=True)
    checkpoint_path = args.output_dir / "iv_rv_0dte_1s_2025_2026_checkpoint.csv"
    jobs = discover_jobs(args.sample_root, args.breeze_root)
    rv_lookup = load_saved_rv_lookup()
    rows = read_checkpoint(checkpoint_path)
    done = {(row["underlying"], row["trade_date"]) for row in rows if row.get("status") == "ok"}
    print(f"jobs={len(jobs)}", flush=True)
    for idx, job in enumerate(jobs, start=1):
        job_key = (job["underlying"], job["trade_date"].isoformat())
        if job_key in done:
            print(f"{idx}/{len(jobs)} {job_key[0]} {job_key[1]} skip checkpoint", flush=True)
            continue
        started = time.perf_counter()
        try:
            row = compute_job(job, args.start_time, args.end_time, args.preload_time, rv_lookup)
            row["status"] = "ok"
        except Exception as exc:
            row = {
                "underlying": job["underlying"],
                "trade_date": job["trade_date"].isoformat(),
                "source": job["source"],
                "status": f"error: {exc}",
            }
        row["seconds"] = round(time.perf_counter() - started, 2)
        rows.append(row)
        pd.DataFrame(rows).to_csv(checkpoint_path, index=False)
        print(
            f"{idx}/{len(jobs)} {row.get('underlying')} {row.get('trade_date')} "
            f"{row.get('status')} iv={row.get('inception_iv')} rv={row.get('rv_day')} "
            f"secs={row['seconds']}",
            flush=True,
        )

    frame = pd.DataFrame(rows)
    csv_path = args.output_dir / "iv_rv_0dte_1s_2025_2026.csv"
    xlsx_path = args.output_dir / "iv_rv_0dte_1s_2025_2026.xlsx"
    frame.to_csv(csv_path, index=False)
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name="iv_rv", index=False)
    print(f"CSV,{csv_path}")
    print(f"XLSX,{xlsx_path}")


def read_checkpoint(path: Path) -> list[dict[str, object]]:
    if not path.exists():
        return []
    frame = pd.read_csv(path)
    return frame.to_dict("records")


def discover_jobs(sample_root: Path, breeze_root: Path) -> list[dict[str, object]]:
    jobs: dict[tuple[str, date], dict[str, object]] = {}
    for underlying in ("NIFTY", "SENSEX"):
        base = sample_root / underlying.lower()
        if not base.exists():
            continue
        for day_dir in base.iterdir():
            if not day_dir.is_dir():
                continue
            try:
                trade_date = datetime.strptime(day_dir.name, "%Y-%m-%d").date()
            except ValueError:
                continue
            if not (date(2025, 1, 1) <= trade_date <= date(2026, 12, 31)):
                continue
            if has_same_day_expiry(day_dir, underlying, trade_date):
                jobs[(underlying, trade_date)] = {
                    "underlying": underlying,
                    "trade_date": trade_date,
                    "source": "sample_hf",
                    "path": sample_root,
                }

    for path in sorted(breeze_root.glob("*.parquet")):
        try:
            meta = pd.read_parquet(path, columns=["trade_date", "expiry_date", "underlying", "product_type"])
        except Exception:
            continue
        meta = meta[meta["product_type"].astype(str).str.lower().eq("options")].copy()
        if meta.empty:
            continue
        meta["trade_date_value"] = pd.to_datetime(meta["trade_date"], errors="coerce").dt.date
        meta["expiry_date_value"] = pd.to_datetime(meta["expiry_date"], errors="coerce").dt.date
        for row in meta[["trade_date_value", "expiry_date_value", "underlying"]].drop_duplicates().itertuples(index=False):
            trade_date = row.trade_date_value
            expiry_date = row.expiry_date_value
            underlying = normalize_underlying(str(row.underlying))
            if trade_date != expiry_date:
                continue
            if not (date(2025, 1, 1) <= trade_date <= date(2026, 12, 31)):
                continue
            key = (underlying, trade_date)
            if key not in jobs:
                jobs[key] = {
                    "underlying": underlying,
                    "trade_date": trade_date,
                    "source": "breeze",
                    "path": path,
                }
    return [jobs[key] for key in sorted(jobs, key=lambda item: (item[1], item[0]))]


def has_same_day_expiry(day_dir: Path, underlying: str, trade_date: date) -> bool:
    option_root = day_dir / "options"
    for path in option_root.rglob("*.csv"):
        match = OPTION_FILE_RE.match(path.name)
        if match is None or match.group("underlying").upper() != underlying:
            continue
        expiry = datetime.strptime(match.group("expiry").upper(), "%d%b%y").date()
        if expiry == trade_date:
            return True
    return False


def compute_job(
    job: dict[str, object],
    start_time: str,
    end_time: str,
    preload_time: str,
    rv_lookup: dict[tuple[str, date], float],
) -> dict[str, object]:
    underlying = str(job["underlying"])
    trade_date = job["trade_date"]
    source = str(job["source"])
    if source == "sample_hf":
        dataset = load_sample_hf_option_dataset(Path(job["path"]), trade_date, underlying)
    else:
        dataset = load_breeze_parquet_option_dataset(Path(job["path"]), trade_date, underlying, preload_time)
    dataset = dataset.__class__(
        frame=dataset.frame[dataset.frame["expiry"].eq(trade_date)].copy(),
        trade_date=dataset.trade_date,
        timestamps=dataset.timestamps,
        underlying=dataset.underlying,
        future_series=dataset.future_series,
    )
    if dataset.frame.empty:
        raise ValueError("no same-day expiry rows after filtering")

    replay = None
    current_time = lambda: replay.now()
    sessions = build_backtest_sessions(dataset, DEFAULT_WORKBOOK, current_time, refresh_ms=0)
    replay = CsvReplayFeed(dataset, sessions)
    register_token_tickers(sessions, dataset)
    if len(sessions) != 1:
        raise ValueError("expected one 0DTE session")
    session = sessions[0]
    surface = dataset.option_surface(session.spec.expiry)
    price_book = DirectPriceBook(dataset, surface, session)

    user_value = session.config.market.user_value
    inception = None
    previous_um = None
    c2c_variance = 0.0
    snapshot_count = 0
    saved_rv = rv_lookup.get((underlying, trade_date))
    while replay.advance():
        timestamp = replay.now()
        hhmm = timestamp.strftime("%H:%M")
        if hhmm < start_time or hhmm > end_time:
            continue
        if hhmm == end_time and timestamp.second > 0:
            continue
        row = surface.loc[dataset.timestamps[replay.position]]
        snapshot = price_book.market_snapshot(row, timestamp, user_value)
        if snapshot is None:
            continue
        user_value = snapshot.user_value

        if previous_um is not None:
            c2c_variance += close_to_close_variance(previous_um, snapshot.universal_mid)
        previous_um = snapshot.universal_mid
        snapshot_count += 1

        if inception is None:
            inception = inception_metrics(session, price_book, row, snapshot, timestamp)
            if inception is not None and saved_rv is not None:
                break

    if inception is None:
        raise ValueError("no clean inception slice")
    rv_day = saved_rv
    if rv_day is None:
        rv_day = scaled_volatility_or_none(
            c2c_variance,
            session.config.market.calendar_days,
            session.config.market.intraday_var,
        )
    iv = inception["inception_iv"] * 100
    return {
        "underlying": underlying,
        "trade_date": trade_date.isoformat(),
        "source": source,
        "inception_timestamp": inception["timestamp"].strftime("%Y-%m-%d %H:%M:%S"),
        "inception_um": inception["universal_mid"],
        "inception_iv": iv,
        "rv_day": rv_day,
        "iv_over_rv": iv - rv_day if rv_day is not None else None,
        "iv_legs": inception["iv_legs"],
        "portfolio_positions": inception["portfolio_positions"],
        "snapshots": snapshot_count,
        "rv_source": "saved_summary" if saved_rv is not None else "raw_recomputed",
    }


def inception_metrics(session, price_book: DirectPriceBook, row: pd.Series, snapshot, timestamp) -> dict[str, object] | None:
    strikes = complete_strikes(price_book, row)
    portfolio = build_light_sample_portfolio(session, snapshot, strikes)
    if portfolio is None:
        return None
    if direct_portfolio_values(portfolio.positions, price_book, row, snapshot, {}) is None:
        return None
    ivs = []
    for position in portfolio.positions:
        price = price_book.option_price(row, position.strike, position.option_type)
        if price is None:
            return None
        try:
            vol = implied_volatility(
                price,
                snapshot.universal_spot,
                position.strike,
                snapshot.time,
                price_book.market.funding_rate,
                position.option_type,
            )
        except (ValueError, ZeroDivisionError, OverflowError):
            return None
        if vol is None or not math.isfinite(vol):
            return None
        ivs.append(vol)
    return {
        "timestamp": timestamp,
        "universal_mid": snapshot.universal_mid,
        "inception_iv": sum(ivs) / len(ivs),
        "iv_legs": len(ivs),
        "portfolio_positions": "; ".join(f"{p.option_type}{p.strike}:{p.lots}" for p in portfolio.positions),
    }


def load_saved_rv_lookup() -> dict[tuple[str, date], float]:
    roots = [
        ROOT / "runs_1s",
        ROOT / "best_synth_sims",
        Path(r"C:\Users\rishi\OneDrive\Desktop\Summary"),
    ]
    candidates = []
    for root in roots:
        if not root.exists():
            continue
        for path in root.rglob("*"):
            if not path.is_file():
                continue
            if path.suffix.lower() == ".csv" and path.name.lower() == "um_mid_bro_summary.csv":
                read_saved_csv_rv(path, candidates)
            elif path.suffix.lower() == ".xlsx" and path.name.lower().endswith("_um_mid_bro_best_market_restrike_diagnostics.xlsx"):
                read_saved_xlsx_rv(path, candidates)
    lookup = {}
    for item in candidates:
        key = (item["underlying"], item["trade_date"])
        if key not in lookup or item["mtime"] > lookup[key]["mtime"]:
            lookup[key] = item
    return {key: item["rv"] for key, item in lookup.items()}


def read_saved_csv_rv(path: Path, candidates: list[dict[str, object]]) -> None:
    parsed = parse_underlying_date(path)
    if parsed is None:
        return
    underlying, trade_date = parsed
    if not (date(2025, 1, 1) <= trade_date <= date(2026, 12, 31)):
        return
    try:
        with path.open(newline="", encoding="utf-8") as handle:
            rows = {row["metric"]: row["value"] for row in csv.DictReader(handle)}
    except Exception:
        return
    rv = numeric(rows.get("c2c_vol") or rows.get("c2c_synth_vol"))
    if rv is not None:
        candidates.append({"underlying": underlying, "trade_date": trade_date, "rv": rv, "mtime": path.stat().st_mtime})


def read_saved_xlsx_rv(path: Path, candidates: list[dict[str, object]]) -> None:
    parsed = parse_underlying_date(path)
    if parsed is None:
        return
    underlying, trade_date = parsed
    if not (date(2025, 1, 1) <= trade_date <= date(2026, 12, 31)):
        return
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if "Summary" not in workbook.sheetnames:
            workbook.close()
            return
        values = {}
        for metric, value in workbook["Summary"].iter_rows(min_row=1, max_col=2, values_only=True):
            if metric is not None:
                values[str(metric).strip()] = value
        workbook.close()
    except Exception:
        return
    rv = numeric(values.get("c2c_vol") or values.get("c2c_synth_vol"))
    if rv is not None:
        candidates.append({"underlying": underlying, "trade_date": trade_date, "rv": rv, "mtime": path.stat().st_mtime})


def parse_underlying_date(path: Path) -> tuple[str, date] | None:
    text = str(path).lower()
    underlying = "NIFTY" if "nifty" in text else "SENSEX" if "sensex" in text else None
    if underlying is None:
        return None
    for pattern, fmt in ((r"_(\d{8})_", "%Y%m%d"), (r"_(\d{2}[a-z]{3}\d{2})_", "%d%b%y")):
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            try:
                return underlying, datetime.strptime(match.group(1), fmt).date()
            except ValueError:
                return None
    return None


def numeric(value) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


if __name__ == "__main__":
    main()
