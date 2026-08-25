from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from backtest.config import DEFAULT_WORKBOOK, normalize_underlying
from backtest.replay import CsvReplayFeed, register_token_tickers
from backtest.sessions import build_backtest_sessions
from backtest_5s.__main__ import (
    DEFAULT_INDEX_ZIP,
    DynamicGammaThresholdPortfolioState,
    default_options_zip,
    load_fyers_dataset,
    parse_date_key,
)
from fit_sensex.ui.app import total_row_numeric_values


PROJECT_ROOT = Path(__file__).resolve().parents[1]
COLUMNS = [
    "timestamp",
    "total_pnl",
    "gamma_lots",
    "options_delta_lots",
    "hedge_delta_lots",
    "total_delta_lots",
]


def main() -> None:
    parser = argparse.ArgumentParser(description="Export focused FYERS 5s PnL/delta workbook.")
    parser.add_argument("--date", required=True)
    parser.add_argument(
        "--underlying",
        required=True,
        type=normalize_underlying,
        choices=("SENSEX", "NIFTY"),
    )
    parser.add_argument("--dynamic-gamma-threshold-ratio", type=float, default=0.40)
    parser.add_argument("--options-zip", type=Path, default=None)
    parser.add_argument("--index-zip", type=Path, default=DEFAULT_INDEX_ZIP)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output", type=Path, default=None)
    args = parser.parse_args()

    trade_date = parse_date_key(args.date)
    options_zip = args.options_zip or default_options_zip(args.underlying, trade_date)
    output_path = args.output or (
        PROJECT_ROOT
        / "outputs"
        / f"{args.underlying.lower()}_5s_{trade_date:%Y%m%d}_pnl_delta.xlsx"
    )

    dataset = load_fyers_dataset(options_zip, args.index_zip, trade_date, args.underlying)
    replay = None
    sessions = build_backtest_sessions(
        dataset,
        args.workbook,
        lambda: replay.now(),
        refresh_ms=0,
    )
    replay = CsvReplayFeed(dataset, sessions)
    register_token_tickers(sessions, dataset)
    session = sessions[0]
    state = DynamicGammaThresholdPortfolioState(session=session)

    rows = []
    while replay.advance():
        timestamp = replay.now()
        result = session.analytics.calculate(session.store.snapshot())
        if result is None:
            continue

        metrics = state.update(
            result,
            timestamp,
            session.config.market.funding_rate,
            session.config.market.brokerage_rate,
            args.dynamic_gamma_threshold_ratio,
        )
        risk_rows = state.portfolio.calculate(result) if state.portfolio else []
        totals = total_row_numeric_values(risk_rows, result) if risk_rows else None
        options_delta_lots = total_value(totals, 19)
        gamma_lots = total_value(totals, 21)
        hedge_delta_lots = state.cumulative_hedge_lots if state.portfolio else None
        rows.append(
            [
                timestamp.replace(tzinfo=None),
                metrics.total_pnl,
                gamma_lots,
                options_delta_lots,
                hedge_delta_lots,
                combined_delta(options_delta_lots, hedge_delta_lots),
            ]
        )

    workbook = build_workbook(rows, args.underlying, trade_date, args.dynamic_gamma_threshold_ratio)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    print(output_path)
    print(len(rows))


def total_value(totals: tuple | None, index: int) -> float | None:
    if totals is None or index >= len(totals):
        return None
    value = totals[index]
    return float(value) if isinstance(value, (int, float)) else None


def combined_delta(options_delta_lots: float | None, hedge_delta_lots: float | None) -> float | None:
    if options_delta_lots is None or hedge_delta_lots is None:
        return None
    return options_delta_lots + hedge_delta_lots


def build_workbook(
    rows: list[list],
    underlying: str,
    trade_date,
    threshold_ratio: float,
) -> Workbook:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    data = workbook.create_sheet("Time Slices")

    data.append(COLUMNS)
    for row in rows:
        data.append(row)
    format_table(data)
    data.freeze_panes = "A2"
    data.auto_filter.ref = data.dimensions

    final = rows[-1] if rows else []
    summary_rows = [
        (f"{underlying} 5s PnL and Delta", None),
        ("Date", trade_date.isoformat()),
        ("Threshold Ratio", threshold_ratio),
        ("Rows", len(rows)),
        ("Final Total PnL", value_by_column(final, "total_pnl")),
        ("Final Gamma Lots", value_by_column(final, "gamma_lots")),
        ("Final Options Delta Lots", value_by_column(final, "options_delta_lots")),
        ("Final Hedge Delta Lots", value_by_column(final, "hedge_delta_lots")),
        ("Final Total Delta Lots", value_by_column(final, "total_delta_lots")),
    ]
    for row in summary_rows:
        summary.append(row)
    summary["A1"].font = Font(size=14, bold=True, color="1F4E78")
    summary.merge_cells("A1:B1")
    for cell in summary["A"]:
        cell.font = Font(bold=True)
    summary["B3"].number_format = "0.00%"
    for row_index in range(5, summary.max_row + 1):
        summary[f"B{row_index}"].number_format = "#,##0.00"
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 18
    return workbook


def value_by_column(row: list, column: str):
    if not row:
        return None
    return row[COLUMNS.index(column)]


def format_table(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    if sheet.max_row > 1:
        table = Table(displayName="Backtest5sPnlDelta", ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
    for row_index in range(2, sheet.max_row + 1):
        sheet[f"A{row_index}"].number_format = "yyyy-mm-dd hh:mm:ss"
        for column in ("B", "C", "D", "E", "F"):
            sheet[f"{column}{row_index}"].number_format = "#,##0.00"
    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
        sheet.column_dimensions[letter].width = min(max(max_length + 2, 13), 24)


if __name__ == "__main__":
    main()
