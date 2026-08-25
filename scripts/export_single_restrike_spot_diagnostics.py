from __future__ import annotations

import argparse
import sys
import time
from dataclasses import replace
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
SCRIPTS = ROOT / "scripts"
for path in (SRC, SCRIPTS):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from backtest.config import DEFAULT_WORKBOOK, normalize_underlying
from backtest.replay import CsvReplayFeed, register_token_tickers
from backtest.sessions import build_backtest_sessions
from backtest_1m.__main__ import dynamic_gamma_threshold
from backtest_sample_hf.__main__ import load_sample_hf_option_dataset, parse_date_key
from export_sample_hf_diagnostics import close_to_close_variance, scaled_volatility_or_none
from export_1s_batch_restrike_workbooks import (
    SegmentInfo,
    in_time_window,
    portfolio_liquidation_premiums,
    portfolio_premiums,
    synthetic_trade_premiums,
)
from run_nifty_1m_light_batch import complete_strikes
from run_sample_hf_pnl_only import (
    DirectHedgeState,
    DirectPriceBook,
    MarketSnapshot,
    build_light_sample_portfolio,
    capture_light_frozen_ivs,
    hedge_multiplier,
)


DEFAULT_DATA_DIR = Path(r"C:\options data\1s data")
DEFAULT_OUTPUT_DIR = ROOT / "10am_restrike_anurag_spot"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--date", required=True)
    parser.add_argument("--underlying", required=True, type=normalize_underlying, choices=("NIFTY", "SENSEX"))
    parser.add_argument("--data-dir", type=Path, default=DEFAULT_DATA_DIR)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--start-time", default="10:00")
    parser.add_argument("--end-time", default="15:15")
    parser.add_argument("--threshold-ratio", type=float, default=0.10)
    args = parser.parse_args()

    trade_date = parse_date_key(args.date)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    threshold_label = f"{args.threshold_ratio:.2f}".replace(".", "p").rstrip("0").rstrip("p")
    output = args.output_dir / (
        f"{args.underlying.lower()}_{trade_date:%d%b%y}_"
        f"{args.start_time.replace(':', '')}_threshold_{threshold_label}_spot_restrike_diagnostics.xlsx"
    )

    total_start = time.perf_counter()
    load_start = time.perf_counter()
    dataset = load_sample_hf_option_dataset(args.data_dir, trade_date, args.underlying)
    load_seconds = time.perf_counter() - load_start
    sim_start = time.perf_counter()
    result = run_spot_path(dataset, DEFAULT_WORKBOOK, args.threshold_ratio, args.start_time, args.end_time)
    sim_seconds = time.perf_counter() - sim_start
    total_seconds = time.perf_counter() - total_start
    write_workbook(
        output,
        args.underlying,
        trade_date,
        args.start_time,
        args.end_time,
        args.threshold_ratio,
        result,
        load_seconds,
        sim_seconds,
        total_seconds,
    )
    summary = result["summary"]
    print(f"OUTPUT,{output}")
    print(f"FINAL_PNL,{summary.get('final_pnl'):.2f}")
    print(f"C2C_VOL,{summary.get('c2c_vol'):.2f}")
    print(f"HEDGE_VOL,{summary.get('hedge_vol'):.2f}")


def run_spot_path(dataset, workbook: Path, threshold_ratio: float, start_time: str, end_time: str):
    if dataset.future_series is None:
        raise ValueError("Spot series is required for the spot-reference run.")
    replay = None
    current_time = lambda: replay.now()
    sessions = build_backtest_sessions(dataset, workbook, current_time, refresh_ms=0)
    replay = CsvReplayFeed(dataset, sessions)
    register_token_tickers(sessions, dataset)
    if len(sessions) != 1:
        raise ValueError("Expected one expiry session.")
    session = sessions[0]
    surface = dataset.option_surface(session.spec.expiry)
    price_book = DirectPriceBook(dataset, surface, session)
    spot_close = spot_close_series(dataset)

    state_info: SegmentInfo | None = None
    waiting_reentry = False
    reentry_done = False
    user_value = session.config.market.user_value
    realized_offset = 0.0
    breach = None
    rows = []
    prev_spot = None
    prev_hedge_spot = None
    c2c_var = 0.0
    hedge_var = 0.0
    net_deltas = []
    hedge_trade_count = 0
    synthetic_buy_lots = 0.0
    synthetic_sell_lots = 0.0
    premium_bought = 0.0
    premium_sold = 0.0
    multiplier = None

    while replay.advance():
        timestamp = replay.now()
        if not in_time_window(timestamp, start_time, end_time):
            continue
        timestamp_value = dataset.timestamps[replay.position]
        if timestamp_value not in spot_close.index or pd.isna(spot_close.loc[timestamp_value]):
            continue
        row = surface.loc[timestamp_value]
        actual_snapshot = price_book.market_snapshot(row, timestamp, user_value)
        if actual_snapshot is None:
            continue
        spot_value = float(spot_close.loc[timestamp_value])
        spot_snapshot = spot_reference_snapshot(actual_snapshot, price_book, spot_value)

        if waiting_reentry:
            state_info = make_spot_segment(session, price_book, row, timestamp, spot_snapshot)
            if state_info is None:
                continue
            multiplier = hedge_multiplier(state_info.state.portfolio.positions)
            bought, sold = portfolio_premiums(state_info.state.portfolio.positions, price_book, row, multiplier)
            premium_bought += bought
            premium_sold += sold
            user_value = state_info.user_value
            waiting_reentry = False

        if state_info is None:
            state_info = make_spot_segment(session, price_book, row, timestamp, spot_snapshot)
            if state_info is None:
                continue
            multiplier = hedge_multiplier(state_info.state.portfolio.positions)
            bought, sold = portfolio_premiums(state_info.state.portfolio.positions, price_book, row, multiplier)
            premium_bought += bought
            premium_sold += sold
            user_value = state_info.user_value

        # Rebuild the spot snapshot with the latest rounded spot user value for the next UM search.
        user_value = spot_snapshot.user_value
        prior_trade_count = len(state_info.state.hedge_trades)
        segment_pnl = state_info.state.update(price_book, row, spot_snapshot, timestamp, threshold_ratio)
        if segment_pnl is None:
            continue

        new_trades = state_info.state.hedge_trades[prior_trade_count:]
        for trade in new_trades:
            lots_change = float(trade["lots_change"])
            if lots_change > 0:
                synthetic_buy_lots += lots_change
            elif lots_change < 0:
                synthetic_sell_lots += abs(lots_change)
            bought, sold = synthetic_trade_premiums(
                price_book,
                row,
                state_info.hedge_strike,
                lots_change,
                multiplier,
            )
            premium_bought += bought
            premium_sold += sold
        hedge_trade_count += len(new_trades)

        running_pnl = realized_offset + segment_pnl
        net_delta = (state_info.state.last_delta_lots or 0.0) + state_info.state.cumulative_hedge_lots
        gamma_lots = state_info.state.last_gamma_lots
        threshold_lots = dynamic_gamma_threshold(gamma_lots, threshold_ratio)
        net_deltas.append(net_delta)

        if prev_spot is not None:
            c2c_var += close_to_close_variance(prev_spot, spot_value)
        prev_spot = spot_value
        traded_spot = spot_value if new_trades else None
        if traded_spot is not None:
            if prev_hedge_spot is not None:
                hedge_var += close_to_close_variance(prev_hedge_spot, traded_spot)
            prev_hedge_spot = traded_spot

        traded_delta_lots = sum(float(trade["lots_change"]) for trade in new_trades)
        rows.append(
            {
                "timestamp": timestamp,
                "universal_mid": actual_snapshot.universal_mid,
                "spot_close": spot_value,
                "running_total_pnl": running_pnl,
                "net_delta_lots": net_delta,
                "gamma_lots": gamma_lots,
                "threshold_lots": threshold_lots,
                "hedge_delta_lots": state_info.state.cumulative_hedge_lots,
                "traded_delta_lots": traded_delta_lots,
            }
        )

        if not reentry_done:
            breach_side = None
            breach_strike = None
            if state_info.short_call is not None and spot_value > state_info.short_call:
                breach_side = "CALL"
                breach_strike = state_info.short_call
            elif state_info.short_put is not None and spot_value < state_info.short_put:
                breach_side = "PUT"
                breach_strike = state_info.short_put
            if breach_side:
                bought, sold = portfolio_liquidation_premiums(
                    state_info.state.portfolio.positions,
                    price_book,
                    row,
                    multiplier,
                )
                premium_bought += bought
                premium_sold += sold
                breach = {
                    "side": breach_side,
                    "timestamp": timestamp,
                    "strike": breach_strike,
                    "spot_close": spot_value,
                    "universal_mid": actual_snapshot.universal_mid,
                    "pnl_at_breach": running_pnl,
                }
                realized_offset = running_pnl
                state_info = None
                waiting_reentry = True
                reentry_done = True

    if not rows:
        raise ValueError("No spot-reference rows produced.")
    summary = {
        "final_pnl": rows[-1]["running_total_pnl"],
        "c2c_vol": scaled_volatility_or_none(c2c_var, session.config.market.calendar_days, session.config.market.intraday_var),
        "hedge_vol": scaled_volatility_or_none(hedge_var, session.config.market.calendar_days, session.config.market.intraday_var),
        "pnl_at_breach": breach["pnl_at_breach"] if breach else None,
        "breach_timestamp": breach["timestamp"].strftime("%Y-%m-%d %H:%M:%S") if breach else None,
        "breach_side": breach["side"] if breach else None,
        "breach_strike": breach["strike"] if breach else None,
        "breach_spot_close": breach["spot_close"] if breach else None,
        "breach_universal_mid": breach["universal_mid"] if breach else None,
        "max_net_delta_lots": max(net_deltas),
        "min_net_delta_lots": min(net_deltas),
        "avg_net_delta_lots": sum(net_deltas) / len(net_deltas),
        "synthetic_hedge_trades": hedge_trade_count,
        "synthetic_buy_lots": synthetic_buy_lots,
        "synthetic_sell_lots": synthetic_sell_lots,
        "synthetic_total_abs_lots": synthetic_buy_lots + synthetic_sell_lots,
        "premium_bought": premium_bought,
        "premium_sold": premium_sold,
    }
    return {"rows": rows, "summary": summary}


def spot_close_series(dataset) -> pd.Series:
    bars = dataset.future_series.bars
    values = pd.Series({pd.Timestamp(ts): bar.close for ts, bar in bars.items()}).sort_index()
    return values.reindex(dataset.timestamps).ffill()


def spot_reference_snapshot(base: MarketSnapshot, price_book: DirectPriceBook, spot_close: float) -> MarketSnapshot:
    next_user_value = round(spot_close / price_book.market.strike_round_base) * price_book.market.strike_round_base
    return replace(
        base,
        universal_mid=spot_close,
        universal_spot=spot_close,
        user_value=next_user_value,
    )


def make_spot_segment(session, price_book: DirectPriceBook, row: pd.Series, timestamp, snapshot: MarketSnapshot) -> SegmentInfo | None:
    strikes = complete_strikes(price_book, row)
    portfolio = build_light_sample_portfolio(session, snapshot, strikes)
    if portfolio is None:
        return None
    frozen_ivs = capture_light_frozen_ivs(portfolio.positions, price_book, row, snapshot)
    if len(frozen_ivs) != len(portfolio.positions):
        return None
    hedge_strike = min(strikes, key=lambda strike: (abs(strike - snapshot.universal_mid), strike))
    short_calls = [position.strike for position in portfolio.positions if position.option_type == "CE" and float(position.lots) < 0]
    short_puts = [position.strike for position in portfolio.positions if position.option_type == "PE" and float(position.lots) < 0]
    return SegmentInfo(
        state=DirectHedgeState(portfolio=portfolio, options_pv_snapshot=None, hedge_strike=hedge_strike),
        short_call=max(short_calls) if short_calls else None,
        short_put=min(short_puts) if short_puts else None,
        hedge_strike=hedge_strike,
        user_value=snapshot.user_value,
        universal_mid=snapshot.universal_mid,
    )


def write_workbook(
    output: Path,
    underlying: str,
    trade_date,
    start_time: str,
    end_time: str,
    threshold_ratio: float,
    result,
    load_seconds: float,
    sim_seconds: float,
    total_seconds: float,
) -> None:
    summary = result["summary"]
    gammas = [row["gamma_lots"] for row in result["rows"] if row.get("gamma_lots") is not None]
    avg_gamma = sum(gammas) / len(gammas) if gammas else None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"{underlying} {trade_date:%d-%b-%y} 10am Spot-Reference Restrike Diagnostics"
    ws["A1"].font = Font(bold=True, size=14)
    rows = [
        ("Underlying", underlying),
        ("Trade Date", trade_date.strftime("%d-%b-%y")),
        ("Reference Price", "Spot Close"),
        ("Start Time", start_time),
        ("End Time", end_time),
        ("Threshold Ratio", threshold_ratio),
        ("Restrike", "Yes"),
        ("Final PnL", summary.get("final_pnl")),
        ("c2c_vol", summary.get("c2c_vol")),
        ("hedge_vol", summary.get("hedge_vol")),
        ("PnL at Breach", summary.get("pnl_at_breach")),
        ("Breach Timestamp", summary.get("breach_timestamp")),
        ("Breach Side", summary.get("breach_side")),
        ("Breach Strike", summary.get("breach_strike")),
        ("Breach Spot Close", summary.get("breach_spot_close")),
        ("Breach Universal Mid", summary.get("breach_universal_mid")),
        ("Max Net Delta Lots", summary.get("max_net_delta_lots")),
        ("Min Net Delta Lots", summary.get("min_net_delta_lots")),
        ("Avg Net Delta Lots", summary.get("avg_net_delta_lots")),
        ("Average Gamma Lots", avg_gamma),
        ("Number of Hedge Events", summary.get("synthetic_hedge_trades")),
        ("Synthetic Buy Lots Hedged", summary.get("synthetic_buy_lots")),
        ("Synthetic Sell Lots Hedged", summary.get("synthetic_sell_lots")),
        ("Total Synthetic Lots Hedged", summary.get("synthetic_total_abs_lots")),
        ("Total Premium Bought", summary.get("premium_bought")),
        ("Total Premium Sold", summary.get("premium_sold")),
        ("Brokerage", brokerage(summary.get("premium_bought"), summary.get("premium_sold"))),
        ("Slippage", slippage(underlying, summary.get("synthetic_total_abs_lots"))),
        ("Data Load Seconds", load_seconds),
        ("Simulation Seconds", sim_seconds),
        ("Total Seconds", total_seconds),
    ]
    for row_idx, (label, value) in enumerate(rows, start=3):
        ws.cell(row_idx, 1).value = label
        ws.cell(row_idx, 1).font = Font(bold=True)
        ws.cell(row_idx, 2).value = value
        if isinstance(value, (int, float)):
            ws.cell(row_idx, 2).number_format = "#,##0" if label in {"Brokerage", "Slippage"} else "#,##0.00"

    path_ws = wb.create_sheet("Restrike Path")
    headers = [
        "timestamp",
        "universal_mid",
        "spot_close",
        "running_total_pnl",
        "net_delta_lots",
        "gamma_lots",
        "threshold_lots",
        "hedge_delta_lots",
        "traded_delta_lots",
    ]
    path_ws.append(headers)
    for cell in path_ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in result["rows"]:
        path_ws.append(
            [
                row["timestamp"].strftime("%Y-%m-%d %H:%M:%S"),
                row["universal_mid"],
                row["spot_close"],
                row["running_total_pnl"],
                row["net_delta_lots"],
                row["gamma_lots"],
                row["threshold_lots"],
                row["hedge_delta_lots"],
                row["traded_delta_lots"],
            ]
        )
    for data_row in path_ws.iter_rows(min_row=2, min_col=2, max_col=9):
        for cell in data_row:
            cell.number_format = "#,##0.00"

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A3" if sheet.title == "Summary" else "A2"
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 20
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def brokerage(premium_bought, premium_sold):
    if premium_bought is None or premium_sold is None:
        return None
    return 0.0024 * ((float(premium_bought) + float(premium_sold)) / 2)


def slippage(underlying: str, synthetic_lots):
    if synthetic_lots is None:
        return None
    normalized = underlying.strip().upper()
    if normalized == "NIFTY":
        return float(synthetic_lots) * 0.125 * 65
    if normalized == "SENSEX":
        return float(synthetic_lots) * 0.41 * 20
    return None


if __name__ == "__main__":
    main()
