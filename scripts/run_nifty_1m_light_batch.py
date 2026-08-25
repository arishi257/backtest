from __future__ import annotations

import argparse
import csv
import io
import json
import os
import re
import shutil
import subprocess
import sys
import time
from datetime import date, datetime
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
SCRIPTS = ROOT / "scripts"
for path in (SRC, SCRIPTS):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from backtest.config import DEFAULT_WORKBOOK
from backtest.data import OptionDataset
from backtest.headless_portfolio import GammaDiffTracker, ParkGammaTracker, is_top_move_time
from backtest.replay import CsvReplayFeed, register_token_tickers
from backtest.sessions import build_backtest_sessions
from backtest_1m.__main__ import dynamic_gamma_threshold
from export_sample_hf_diagnostics import (
    close_to_close_variance,
    round_or_none,
    scaled_volatility_or_none,
)
from run_sample_hf_pnl_only import (
    DirectHedgeState,
    DirectPriceBook,
    build_light_sample_portfolio,
    capture_light_frozen_ivs,
)


DATA_DIR = Path(r"C:\options data\NIFTY")
OUTPUT_WORKBOOK = ROOT / "runs_1m" / "nifty_0dte_dates_2019_2026.xlsx"
DIAGNOSTICS_DIR = ROOT / "runs_1m" / "10am_diagnostics"
CACHE_DIR = ROOT / ".backtest_data_cache" / "nifty_1m_light"
NODE_EXE = Path(r"C:\Users\rishi\.cache\codex-runtimes\codex-primary-runtime\dependencies\node\bin\node.exe")
DATE_RE = re.compile(r"(?P<date>\d{8})")
TICKER_RE = re.compile(r"^NIFTY(?P<expiry>\d{2}[A-Z]{3}\d{2})(?P<strike>\d+)(?P<option_type>CE|PE)\.NFO$")
IST = "Asia/Kolkata"
START_TIME = "10:00"


def main() -> None:
    parser = argparse.ArgumentParser(description="Run NIFTY 1m light simulations for 2026 0DTE dates.")
    parser.add_argument("--data-dir", type=Path, default=DATA_DIR)
    parser.add_argument("--workbook", type=Path, default=OUTPUT_WORKBOOK)
    parser.add_argument("--holidays", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--diagnostics-dir", type=Path, default=DIAGNOSTICS_DIR)
    parser.add_argument("--threshold-ratio", type=float, default=0.40)
    parser.add_argument("--start-time", default=START_TIME)
    parser.add_argument("--year", default="2026")
    args = parser.parse_args()

    dates = load_year_dates(args.workbook, args.year)
    results = []
    args.diagnostics_dir.mkdir(parents=True, exist_ok=True)
    for trade_date in dates:
        start = time.perf_counter()
        print(f"Running NIFTY {trade_date:%d-%b-%Y}...")
        try:
            dataset, cache_used = load_nifty_1m_dataset(args.data_dir, trade_date)
            diagnostics, portfolio_rows, summary, cycles = run_light_1m(
                dataset,
                args.holidays,
                args.threshold_ratio,
                args.start_time,
            )
            run_dir = ROOT / "runs_1m" / f"_tmp_nifty_1m_{trade_date:%Y%m%d}_{datetime.now():%Y%m%d_%H%M%S}"
            run_dir.mkdir(parents=True, exist_ok=True)
            write_csv(run_dir / "diagnostics.csv", diagnostics)
            write_csv(run_dir / "portfolio_0920.csv", portfolio_rows)
            write_summary_csv(run_dir / "summary_metrics.csv", summary)
            temp_workbook = build_diagnostics_workbook(run_dir, "NIFTY", trade_date, args.start_time)
            final_workbook = unique_output_path(args.diagnostics_dir, trade_date)
            shutil.move(str(temp_workbook), final_workbook)
            shutil.rmtree(run_dir, ignore_errors=True)
            elapsed = time.perf_counter() - start
            results.append(
                {
                    "date": trade_date,
                    "portfolio_total_pnl": summary.get("portfolio_total_pnl"),
                    "frozen_iv_total_pnl": summary.get("frozen_iv_total_pnl"),
                    "c2c_synth_vol": summary.get("c2c_synth_vol"),
                    "hedge_vol": summary.get("hedge_vol"),
                    "um_start": diagnostics[0].get("universal_mid") if diagnostics else None,
                    "um_end": diagnostics[-1].get("universal_mid") if diagnostics else None,
                    "workbook": final_workbook,
                    "status": "ok",
                    "cache_used": cache_used,
                    "seconds": elapsed,
                    "cycles": cycles,
                }
            )
            print(
                f"  total={format_number(summary.get('portfolio_total_pnl'))}, "
                f"frozen={format_number(summary.get('frozen_iv_total_pnl'))}, "
                f"cache={'yes' if cache_used else 'no'}, seconds={elapsed:.2f}"
            )
        except Exception as exc:  # noqa: BLE001 - batch should continue.
            elapsed = time.perf_counter() - start
            results.append(
                {
                    "date": trade_date,
                    "portfolio_total_pnl": None,
                    "frozen_iv_total_pnl": None,
                    "c2c_synth_vol": None,
                    "hedge_vol": None,
                    "um_start": None,
                    "um_end": None,
                    "workbook": None,
                    "status": f"error: {exc}",
                    "cache_used": False,
                    "seconds": elapsed,
                    "cycles": 0,
                }
            )
            print(f"  ERROR: {exc}")

    update_results_workbook(args.workbook, args.year, results)
    print(f"Updated {args.workbook}")
    print("date,portfolio_total_pnl,frozen_iv_total_pnl,status,workbook")
    for row in results:
        print(
            f"{row['date']:%d-%b-%y},{format_number(row['portfolio_total_pnl'])},"
            f"{format_number(row['frozen_iv_total_pnl'])},{row['status']},{row['workbook'] or ''}"
        )


def load_year_dates(workbook_path: Path, sheet_name: str) -> list[date]:
    workbook = openpyxl.load_workbook(workbook_path, data_only=True)
    sheet = workbook[sheet_name]
    dates = []
    for row in range(4, sheet.max_row + 1):
        value = sheet.cell(row, 1).value
        if value is None or str(value).startswith("Count"):
            continue
        if isinstance(value, datetime):
            dates.append(value.date())
        else:
            dates.append(datetime.strptime(str(value), "%d-%b-%y").date())
    workbook.close()
    return dates


def load_nifty_1m_dataset(data_dir: Path, trade_date: date) -> tuple[OptionDataset, bool]:
    cache_path = cache_file_path(data_dir, trade_date)
    if cache_path.exists():
        cached = pd.read_pickle(cache_path)
        return OptionDataset(
            frame=cached["frame"],
            trade_date=trade_date,
            timestamps=cached["timestamps"],
            underlying="NIFTY",
            future_series=None,
        ), True
    source = find_daily_source(data_dir, trade_date)
    if source is None:
        raise FileNotFoundError(f"No NIFTY 1m source found for {trade_date:%d-%b-%Y}.")
    raw = read_daily_csv(source)
    frame = parse_option_frame(raw, trade_date)
    timestamps = pd.date_range(
        pd.Timestamp(f"{trade_date} 09:15:00", tz=IST),
        pd.Timestamp(f"{trade_date} 15:24:00", tz=IST),
        freq="min",
    )
    cached = {"frame": frame, "timestamps": timestamps}
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    pd.to_pickle(cached, cache_path)
    return OptionDataset(frame, trade_date, timestamps, "NIFTY", None), False


def find_daily_source(data_dir: Path, trade_date: date) -> tuple[Path, str | None, str | None] | None:
    year_dir = data_dir / f"{trade_date:%Y}"
    date_text = f"{trade_date:%d%m%Y}"
    if not year_dir.exists():
        return None
    for csv_path in sorted(year_dir.rglob(f"*{date_text}*.csv")):
        return (csv_path, None, None)
    for zip_path in sorted(year_dir.rglob("*.zip")):
        try:
            with ZipFile(zip_path) as zipped:
                for member in sorted(zipped.namelist()):
                    if date_text not in Path(member).name:
                        continue
                    if member.lower().endswith(".csv"):
                        return (zip_path, member, None)
                    if member.lower().endswith(".zip"):
                        with ZipFile(io.BytesIO(zipped.read(member))) as nested:
                            for nested_member in sorted(nested.namelist()):
                                if nested_member.lower().endswith(".csv") and date_text in Path(nested_member).name:
                                    return (zip_path, member, nested_member)
        except BadZipFile:
            continue
    return None


def read_daily_csv(source: tuple[Path, str | None, str | None]) -> pd.DataFrame:
    path, member, nested_member = source
    columns = ["Ticker", "Date", "Time", "Close"]
    if member is None:
        return pd.read_csv(path, usecols=columns)
    with ZipFile(path) as outer:
        if nested_member is None:
            with outer.open(member) as handle:
                return pd.read_csv(handle, usecols=columns)
        nested_bytes = outer.read(member)
        with ZipFile(io.BytesIO(nested_bytes)) as inner:
            with inner.open(nested_member) as handle:
                return pd.read_csv(handle, usecols=columns)


def parse_option_frame(raw: pd.DataFrame, trade_date: date) -> pd.DataFrame:
    tickers = raw["Ticker"].astype(str).str.upper()
    parsed = tickers.str.extract(TICKER_RE)
    data = raw[parsed["expiry"].notna()].copy()
    parsed = parsed.loc[data.index]
    expiry = pd.to_datetime(parsed["expiry"], format="%d%b%y", errors="coerce").dt.date
    data["expiry"] = expiry
    data = data[data["expiry"].eq(trade_date)]
    if data.empty:
        raise ValueError(f"No NIFTY 0DTE option rows found for {trade_date:%d-%b-%Y}.")
    data["ticker"] = tickers.loc[data.index]
    data["strike"] = parsed.loc[data.index, "strike"].astype(int)
    data["option_type"] = parsed.loc[data.index, "option_type"].str.upper()
    data["close"] = pd.to_numeric(data["Close"], errors="coerce")
    timestamp_text = data["Date"].astype(str) + " " + data["Time"].astype(str)
    data["timestamp"] = pd.to_datetime(timestamp_text, dayfirst=True, errors="coerce").dt.tz_localize(IST).dt.floor("min")
    data = data.dropna(subset=["timestamp", "close"])
    return data[["ticker", "timestamp", "expiry", "strike", "option_type", "close"]].sort_values(
        ["timestamp", "ticker"]
    ).reset_index(drop=True)


def run_light_1m(dataset, workbook: Path, threshold_ratio: float, start_time: str):
    replay = None
    current_time = lambda: replay.now()
    sessions = build_backtest_sessions(dataset, workbook, current_time, refresh_ms=0)
    replay = CsvReplayFeed(dataset, sessions)
    register_token_tickers(sessions, dataset)
    if len(sessions) != 1:
        raise ValueError("Expected one expiry session.")
    session = sessions[0]
    surface = dataset.option_surface(session.spec.expiry)
    price_book = DirectPriceBook(dataset, surface, session)

    market_state = None
    frozen_state = None
    portfolio_rows = []
    diagnostics = []
    universal_mid_points = []
    gamma_l_points = []
    park_gamma_tracker = ParkGammaTracker()
    previous_universal_mid = None
    previous_hedge_universal_mid = None
    c2c_synth_variance = 0.0
    hedge_variance = 0.0
    calendar_days = session.config.market.calendar_days
    intraday_var = session.config.market.intraday_var
    user_value = session.config.market.user_value
    summary = {}
    cycles = 0

    while replay.advance():
        timestamp = replay.now()
        if timestamp.strftime("%H:%M") < start_time:
            continue
        cycles += 1
        timestamp_value = dataset.timestamps[replay.position]
        row = surface.loc[timestamp_value]
        snapshot = price_book.market_snapshot(row, timestamp, user_value)
        if snapshot is None:
            continue
        user_value = snapshot.user_value
        if market_state is None or frozen_state is None:
            strikes = complete_strikes(price_book, row)
            portfolio = build_light_sample_portfolio(session, snapshot, strikes)
            if portfolio is None:
                continue
            frozen_ivs = capture_light_frozen_ivs(portfolio.positions, price_book, row, snapshot)
            if len(frozen_ivs) != len(portfolio.positions):
                continue
            hedge_strike = min(strikes, key=lambda strike: (abs(strike - snapshot.universal_mid), strike))
            market_state = DirectHedgeState(portfolio=portfolio, options_pv_snapshot=None, hedge_strike=hedge_strike)
            frozen_state = DirectHedgeState(
                portfolio=portfolio,
                options_pv_snapshot=None,
                hedge_strike=hedge_strike,
                frozen_ivs=frozen_ivs,
            )
            portfolio_rows = portfolio_snapshot_rows(timestamp, snapshot.universal_mid, portfolio.positions)

        prior_trade_count = len(market_state.hedge_trades)
        prior_frozen_trade_count = len(frozen_state.hedge_trades)
        market_pnl = market_state.update(price_book, row, snapshot, timestamp, threshold_ratio)
        frozen_pnl = frozen_state.update(price_book, row, snapshot, timestamp, threshold_ratio)
        if market_pnl is None:
            continue

        gamma_lots = market_state.last_gamma_lots
        threshold_lots = dynamic_gamma_threshold(gamma_lots, threshold_ratio)
        net_delta_lots = market_state.last_delta_lots + market_state.cumulative_hedge_lots
        traded_delta_lots = sum(float(trade["lots_change"]) for trade in market_state.hedge_trades[prior_trade_count:])
        traded_universal_mid = snapshot.universal_mid if abs(traded_delta_lots) > 1e-9 else None
        frozen_threshold_lots = dynamic_gamma_threshold(frozen_state.last_gamma_lots, threshold_ratio)
        frozen_net_delta_lots = frozen_state.last_delta_lots + frozen_state.cumulative_hedge_lots
        frozen_traded_delta_lots = sum(
            float(trade["lots_change"]) for trade in frozen_state.hedge_trades[prior_frozen_trade_count:]
        )
        frozen_traded_universal_mid = snapshot.universal_mid if abs(frozen_traded_delta_lots) > 1e-9 else None

        if previous_universal_mid is not None and is_top_move_time(timestamp):
            c2c_synth_variance += close_to_close_variance(previous_universal_mid, snapshot.universal_mid)
        if is_top_move_time(timestamp):
            previous_universal_mid = snapshot.universal_mid
        if traded_universal_mid is not None:
            if previous_hedge_universal_mid is not None:
                hedge_variance += close_to_close_variance(previous_hedge_universal_mid, traded_universal_mid)
            previous_hedge_universal_mid = traded_universal_mid

        if gamma_lots is not None:
            gamma_l_points.append((timestamp, gamma_lots))
        universal_mid_points.append((timestamp, snapshot.universal_mid))
        park_gamma_metrics = park_gamma_tracker.update(None, gamma_lots)

        diagnostics.append(
            {
                "timestamp": timestamp.strftime("%Y-%m-%d %H:%M:%S"),
                "universal_mid": round_or_none(snapshot.universal_mid),
                "running_total_pnl": round_or_none(market_pnl),
                "gamma_lots": round_or_none(gamma_lots),
                "threshold_lots": round_or_none(threshold_lots),
                "net_delta_lots_options_plus_hedge": round_or_none(net_delta_lots),
                "hedge_delta_lots": round_or_none(market_state.cumulative_hedge_lots),
                "traded_delta_lots": round_or_none(traded_delta_lots),
                "traded_universal_mid": round_or_none(traded_universal_mid),
                "frozen_iv_running_total_pnl": round_or_none(frozen_pnl),
                "frozen_iv_gamma_lots": round_or_none(frozen_state.last_gamma_lots),
                "frozen_iv_threshold_lots": round_or_none(frozen_threshold_lots),
                "frozen_iv_net_delta_lots_options_plus_hedge": round_or_none(frozen_net_delta_lots),
                "frozen_iv_hedge_delta_lots": round_or_none(frozen_state.cumulative_hedge_lots),
                "frozen_iv_traded_delta_lots": round_or_none(frozen_traded_delta_lots),
                "frozen_iv_traded_universal_mid": round_or_none(frozen_traded_universal_mid),
            }
        )
        summary = {
            "portfolio_total_pnl": market_pnl,
            "portfolio_gamma_l": gamma_lots,
            "portfolio_gamma_diff_total": None,
            "park_gamma_pnl_diff_total": park_gamma_metrics.park_gamma_pnl_diff_total,
            "gk_gamma_pnl_diff_total": park_gamma_metrics.gk_gamma_pnl_diff_total,
            "frozen_iv_total_pnl": frozen_pnl,
        }

    if not diagnostics:
        raise ValueError("No 1m light diagnostics rows were produced.")
    if universal_mid_points and gamma_l_points:
        tracker = GammaDiffTracker()
        tracker.universal_mid_points = universal_mid_points
        tracker.gamma_l_points = gamma_l_points
        summary["portfolio_gamma_diff_total"] = tracker.total()
    summary["c2c_spot_vol"] = None
    summary["park_vol"] = None
    summary["gk_vol"] = None
    summary["c2c_synth_vol"] = scaled_volatility_or_none(c2c_synth_variance, calendar_days, intraday_var)
    summary["hedge_vol"] = scaled_volatility_or_none(hedge_variance, calendar_days, intraday_var)
    return diagnostics, portfolio_rows, summary, cycles


def complete_strikes(price_book: DirectPriceBook, row) -> list[int]:
    return [
        strike
        for strike in price_book.strikes
        if price_book.option_price(row, strike, "CE") is not None
        and price_book.option_price(row, strike, "PE") is not None
    ]


def portfolio_snapshot_rows(timestamp, universal_mid, positions) -> list[dict[str, object]]:
    return [
        {
            "timestamp": timestamp.strftime("%Y-%m-%d %H:%M:%S"),
            "universal_mid": round_or_none(universal_mid),
            "underlying": position.underlying,
            "maturity": position.maturity,
            "strike": position.strike,
            "option_type": position.option_type,
            "lots": position.lots,
            "qty": position.qty,
            "mult": position.mult,
        }
        for position in positions
    ]


def build_diagnostics_workbook(run_dir: Path, underlying: str, trade_date: date, start_time: str) -> Path:
    env = os.environ.copy()
    env["RUN_DIR"] = str(run_dir)
    env["UNDERLYING"] = underlying
    env["TRADE_DATE"] = trade_date.isoformat()
    env["START_TIME_LABEL"] = start_time
    env["INTERVAL_LABEL"] = "1m"
    env["CLEAN_INPUTS"] = "1"
    env["NODE_PATH"] = r"C:\Users\rishi\.cache\codex-runtimes\codex-primary-runtime\dependencies\node\node_modules"
    subprocess.run(
        [str(NODE_EXE), str(ROOT / "scripts" / "build_sample_hf_diagnostics_workbook.mjs")],
        cwd=ROOT,
        env=env,
        check=True,
    )
    workbooks = sorted(run_dir.glob("*.xlsx"))
    if not workbooks:
        raise RuntimeError(f"No workbook created in {run_dir}.")
    return workbooks[-1]


def update_results_workbook(path: Path, sheet_name: str, results: list[dict[str, object]]) -> None:
    workbook = openpyxl.load_workbook(path)
    sheet = workbook[sheet_name]
    headers = [
        "Date",
        "Total PnL",
        "Frozen IV PnL",
        "c2c_synth_vol",
        "hedge_vol",
        "UM mid @ 10.00",
        "UM mid @ 15.24",
        "Diagnostics Workbook",
        "Status",
    ]
    for column, header in enumerate(headers, start=1):
        sheet.cell(3, column).value = header
    by_date = {row["date"]: row for row in results}
    for row_index in range(4, sheet.max_row + 1):
        value = sheet.cell(row_index, 1).value
        if value is None or str(value).startswith("Count"):
            continue
        trade_date = datetime.strptime(str(value), "%d-%b-%y").date()
        result = by_date.get(trade_date)
        if result is None:
            continue
        sheet.cell(row_index, 2).value = round_or_none(result["portfolio_total_pnl"])
        sheet.cell(row_index, 3).value = round_or_none(result["frozen_iv_total_pnl"])
        sheet.cell(row_index, 4).value = round_or_none(result["c2c_synth_vol"])
        sheet.cell(row_index, 5).value = round_or_none(result["hedge_vol"])
        sheet.cell(row_index, 6).value = result["um_start"]
        sheet.cell(row_index, 7).value = result["um_end"]
        sheet.cell(row_index, 8).value = str(result["workbook"]) if result["workbook"] else None
        sheet.cell(row_index, 9).value = result["status"]
    for column in ("B", "C", "D", "E", "F", "G"):
        for cell in sheet[column]:
            cell.number_format = "#,##0.00"
    for letter, width in {
        "A": 14,
        "B": 14,
        "C": 16,
        "D": 15,
        "E": 13,
        "F": 16,
        "G": 16,
        "H": 95,
        "I": 24,
    }.items():
        sheet.column_dimensions[letter].width = width
    workbook.save(path)


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def write_summary_csv(path: Path, summary: dict[str, object]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=["metric", "value"])
        writer.writeheader()
        for key in (
            "portfolio_total_pnl",
            "portfolio_gamma_l",
            "portfolio_gamma_diff_total",
            "park_gamma_pnl_diff_total",
            "gk_gamma_pnl_diff_total",
            "frozen_iv_total_pnl",
            "c2c_spot_vol",
            "park_vol",
            "gk_vol",
            "c2c_synth_vol",
            "hedge_vol",
        ):
            writer.writerow({"metric": key, "value": round_or_none(summary.get(key))})


def unique_output_path(output_dir: Path, trade_date: date) -> Path:
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    date_name = f"{trade_date.day:02d}{months[trade_date.month - 1]}{trade_date.year % 100:02d}"
    base = f"nifty_{date_name}_1m_10am_light_diagnostics"
    path = output_dir / f"{base}.xlsx"
    suffix = 2
    while path.exists():
        path = output_dir / f"{base}_{suffix}.xlsx"
        suffix += 1
    return path


def cache_file_path(data_dir: Path, trade_date: date) -> Path:
    key = re.sub(r"[^A-Za-z0-9_.-]+", "_", str(data_dir.resolve()))
    return CACHE_DIR / f"{key}_{trade_date:%Y%m%d}.pkl"


def format_number(value: object) -> str:
    return "--" if value is None else f"{float(value):.2f}"


if __name__ == "__main__":
    main()
