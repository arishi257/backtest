from __future__ import annotations

import argparse
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
PYTHON = ROOT.parent / ".venv" / "Scripts" / "python.exe"
DEFAULT_DATA_DIR = Path(r"C:\options data\1s data")
DEFAULT_OUTPUT_DIR = ROOT / "best_synth_sims"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dates", nargs="+", required=True, help="Date labels like 30Dec25.")
    parser.add_argument("--data-dir", type=Path, default=DEFAULT_DATA_DIR)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--start-time", default="10:00")
    parser.add_argument("--end-time", default="15:15")
    parser.add_argument("--threshold-ratio", default="0.10")
    args = parser.parse_args()

    args.output_dir.mkdir(parents=True, exist_ok=True)
    all_rows = []
    for date_label in args.dates:
        trade_date = parse_date_label(date_label)
        underlying = detect_underlying(args.data_dir, trade_date)
        print(f"RUNNING,{date_label},{underlying}", flush=True)
        rows = run_date(
            trade_date,
            date_label,
            underlying,
            args.data_dir,
            args.output_dir,
            args.start_time,
            args.end_time,
            args.threshold_ratio,
        )
        all_rows.extend(rows)
        for row in rows:
            print(
                "RESULT,"
                + ",".join(
                    [
                        row["date"],
                        row["underlying"],
                        row["version"],
                        f"{row['final_pnl']:.2f}",
                        str(row["hedge_events"]),
                        str(row["synth_lots"]),
                        f"{row['run_seconds']:.2f}",
                    ]
                ),
                flush=True,
            )

    print("TABLE")
    for row in all_rows:
        print(
            "| "
            + " | ".join(
                [
                    row["date"],
                    row["underlying"],
                    row["version"],
                    f"{row['final_pnl']:,.0f}",
                    f"{row['hedge_events']:,.0f}",
                    f"{row['synth_lots']:,.0f}",
                    f"{row['premium_bought']:,.0f}",
                    f"{row['premium_sold']:,.0f}",
                    f"{row['brokerage']:,.0f}",
                    f"{row['slippage']:,.0f}",
                    f"{row['run_seconds']:,.2f}",
                ]
            )
            + " |"
        )


def run_date(
    trade_date,
    date_label: str,
    underlying: str,
    data_dir: Path,
    output_dir: Path,
    start_time: str,
    end_time: str,
    threshold_ratio: str,
) -> list[dict[str, object]]:
    date_arg = trade_date.strftime("%d%m%Y")
    prefix = f"{underlying.lower()}_{date_label}_1000_threshold_0p1"
    jobs = [
        (
            "Regular",
            [
                str(PYTHON),
                "scripts\\export_single_restrike_diagnostics.py",
                "--date",
                date_arg,
                "--underlying",
                underlying,
                "--data-dir",
                str(data_dir),
                "--output-dir",
                str(output_dir),
                "--start-time",
                start_time,
                "--end-time",
                end_time,
                "--threshold-ratio",
                threshold_ratio,
            ],
            output_dir / f"{prefix}_restrike_diagnostics.xlsx",
        ),
        (
            "Cheapest Synth",
            [
                str(PYTHON),
                "scripts\\export_single_restrike_cheapest_synth_diagnostics.py",
                "--date",
                date_arg,
                "--underlying",
                underlying,
                "--data-dir",
                str(data_dir),
                "--output-dir",
                str(output_dir),
                "--start-time",
                start_time,
                "--end-time",
                end_time,
                "--threshold-ratio",
                threshold_ratio,
                "--hedge-strike-rule",
                "cheapest",
            ],
            output_dir / f"{prefix}_cheapest_synth_restrike_diagnostics.xlsx",
        ),
    ]

    timings = {}
    for version, command, _ in jobs:
        start = time.perf_counter()
        process = subprocess.run(command, cwd=ROOT, text=True, capture_output=True)
        timings[version] = time.perf_counter() - start
        if process.returncode != 0:
            sys.stdout.write(process.stdout)
            sys.stderr.write(process.stderr)
            raise RuntimeError(f"{date_label} {underlying} {version} failed.")

    remove_other_versions(output_dir, prefix)
    rows = []
    for version, _, path in jobs:
        summary = read_summary(path)
        rows.append(
            {
                "date": format_hyphen_date(date_label),
                "underlying": underlying,
                "version": version,
                "final_pnl": float(summary.get("Final PnL") or 0.0),
                "hedge_events": float(summary.get("Number of Hedge Events") or 0.0),
                "synth_lots": float(summary.get("Total Synthetic Lots Hedged") or 0.0),
                "premium_bought": float(summary.get("Total Premium Bought") or 0.0),
                "premium_sold": float(summary.get("Total Premium Sold") or 0.0),
                "brokerage": float(summary.get("Brokerage") or 0.0),
                "slippage": float(summary.get("Slippage") or 0.0),
                "run_seconds": timings[version],
            }
        )
    write_table_tabs(jobs, f"{underlying} {format_hyphen_date(date_label)} 10am Restrike Comparison", rows)
    return rows


def detect_underlying(data_dir: Path, trade_date) -> str:
    token = trade_date.strftime("%d%b%y").upper()
    found = []
    for underlying in ("NIFTY", "SENSEX"):
        option_dir = data_dir / underlying.lower() / trade_date.strftime("%Y-%m-%d") / "options"
        if not option_dir.exists():
            continue
        for path in option_dir.rglob("*.csv"):
            if token in path.name.upper():
                found.append(underlying)
                break
    if not found:
        raise FileNotFoundError(f"No same-day expiry option data found for {trade_date:%d-%b-%y}.")
    return found[0]


def read_summary(path: Path) -> dict[str, object]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook["Summary"]
    data = {}
    for row in worksheet.iter_rows(values_only=True):
        if row and row[0] is not None:
            data[str(row[0])] = row[1]
    workbook.close()
    return data


def write_table_tabs(jobs, title: str, rows: list[dict[str, object]]) -> None:
    headers = [
        "Version",
        "Final PnL",
        "Hedge Events",
        "Synth Lots",
        "Premium Bought",
        "Premium Sold",
        "Brokerage",
        "Slippage",
        "Run Seconds",
    ]
    values = [
        [
            row["version"],
            row["final_pnl"],
            row["hedge_events"],
            row["synth_lots"],
            row["premium_bought"],
            row["premium_sold"],
            row["brokerage"],
            row["slippage"],
            row["run_seconds"],
        ]
        for row in rows
    ]
    for _, _, path in jobs:
        workbook = load_workbook(path)
        if "table" in workbook.sheetnames:
            del workbook["table"]
        worksheet = workbook.create_sheet("table", 0)
        worksheet["A1"] = title
        worksheet["A1"].font = Font(bold=True, size=14)
        worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        worksheet.append([])
        worksheet.append(headers)
        header_fill = PatternFill("solid", fgColor="1F4E78")
        subtle_fill = PatternFill("solid", fgColor="D9EAF7")
        white_font = Font(color="FFFFFF", bold=True)
        bold = Font(bold=True)
        thin = Side(style="thin", color="B7B7B7")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for cell in worksheet[3]:
            cell.fill = header_fill
            cell.font = white_font
            cell.alignment = Alignment(horizontal="center")
            cell.border = border
        for row in values:
            worksheet.append(row)
        for row in worksheet.iter_rows(min_row=4, max_row=3 + len(values), min_col=1, max_col=len(headers)):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(horizontal="right" if cell.column > 1 else "left")
            row[0].font = bold
        for col in range(2, len(headers) + 1):
            for col_cells in worksheet.iter_cols(min_col=col, max_col=col, min_row=4, max_row=3 + len(values)):
                for cell in col_cells:
                    cell.number_format = "#,##0.00" if col == len(headers) else "#,##0"
        for cell in worksheet[4]:
            cell.fill = subtle_fill
        worksheet.freeze_panes = "A4"
        widths = [24, 14, 14, 14, 18, 18, 14, 14, 14]
        for index, width in enumerate(widths, start=1):
            worksheet.column_dimensions[get_column_letter(index)].width = width
        workbook.save(path)
        workbook.close()


def remove_other_versions(output_dir: Path, prefix: str) -> None:
    for suffix in (
        "regular_delayed_synth",
        "cheapest_delayed_synth",
        "closest_um_synth",
        "best_market_synth",
        "best_market_delayed_synth",
    ):
        path = output_dir / f"{prefix}_{suffix}_restrike_diagnostics.xlsx"
        if path.exists():
            path.unlink()


def parse_date_label(value: str):
    return datetime.strptime(value, "%d%b%y").date()


def format_hyphen_date(value: str) -> str:
    return f"{value[:2]}-{value[2:5]}-{value[5:]}"


if __name__ == "__main__":
    main()
