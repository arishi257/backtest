from __future__ import annotations

import argparse
import sys
import time
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
SCRIPTS = ROOT / "scripts"
for path in (SRC, SCRIPTS):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from backtest.config import DEFAULT_WORKBOOK, normalize_underlying
from backtest_sample_hf.__main__ import load_sample_hf_option_dataset, parse_date_key
from export_1s_batch_restrike_workbooks import run_path


DEFAULT_DATA_DIR = Path(r"C:\options data\1s data")
DEFAULT_OUTPUT_DIR = Path(r"C:\Users\rishi\OneDrive\Desktop\Summary")
START_TIMES = ("09:20", "10:00")
THRESHOLDS = (0.1, 0.4, 0.9)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--date", required=True)
    parser.add_argument("--underlying", required=True, type=normalize_underlying, choices=("NIFTY", "SENSEX"))
    parser.add_argument("--data-dir", type=Path, default=DEFAULT_DATA_DIR)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--end-time", default="15:15")
    args = parser.parse_args()

    trade_date = parse_date_key(args.date)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    output = args.output_dir / f"{args.underlying.lower()}_{trade_date:%d%b%y}_1s_light_scenario_grid.xlsx"

    total_start = time.perf_counter()
    load_start = time.perf_counter()
    dataset = load_sample_hf_option_dataset(args.data_dir, trade_date, args.underlying)
    load_seconds = time.perf_counter() - load_start

    results = []
    for start_time in START_TIMES:
        for threshold in THRESHOLDS:
            for restrike in (False, True):
                scenario_start = time.perf_counter()
                result = run_path(
                    dataset,
                    DEFAULT_WORKBOOK,
                    threshold,
                    start_time,
                    args.end_time,
                    with_restrike=restrike,
                )
                results.append(
                    {
                        "start_time": start_time,
                        "threshold": threshold,
                        "restrike": "Yes" if restrike else "No",
                        "final_pnl": result["summary"]["final_pnl"],
                        "seconds": time.perf_counter() - scenario_start,
                    }
                )

    total_seconds = time.perf_counter() - total_start
    write_workbook(output, args.underlying, trade_date, args.end_time, load_seconds, total_seconds, results)

    print(f"OUTPUT,{output}")
    print(f"TOTAL_SECONDS,{total_seconds:.2f}")
    print("START,THRESHOLD,NO_RESTRIKE,RESTRIKE")
    for start_time in START_TIMES:
        for threshold in THRESHOLDS:
            no_restrike = next(
                row["final_pnl"]
                for row in results
                if row["start_time"] == start_time and row["threshold"] == threshold and row["restrike"] == "No"
            )
            restrike = next(
                row["final_pnl"]
                for row in results
                if row["start_time"] == start_time and row["threshold"] == threshold and row["restrike"] == "Yes"
            )
            print(f"{start_time},{threshold:.1f},{no_restrike:.2f},{restrike:.2f}")


def write_workbook(
    output: Path,
    underlying: str,
    trade_date: date,
    end_time: str,
    load_seconds: float,
    total_seconds: float,
    results: list[dict[str, object]],
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scenario PnL"
    ws["A1"] = f"{underlying} {trade_date:%d-%b-%y} 1s Light Scenario Grid"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "End Time"
    ws["B2"] = end_time
    ws["A3"] = "Data Load Seconds"
    ws["B3"] = load_seconds
    ws["A4"] = "Total Seconds"
    ws["B4"] = total_seconds

    headers = ["Start Time", "Threshold", "Restrike", "Final PnL", "Scenario Seconds"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(6, col)
        cell.value = header
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row_idx, row in enumerate(results, start=7):
        ws.cell(row_idx, 1).value = row["start_time"]
        ws.cell(row_idx, 2).value = row["threshold"]
        ws.cell(row_idx, 3).value = row["restrike"]
        ws.cell(row_idx, 4).value = row["final_pnl"]
        ws.cell(row_idx, 5).value = row["seconds"]
        ws.cell(row_idx, 4).number_format = "#,##0.00"
        ws.cell(row_idx, 5).number_format = "#,##0.00"

    matrix = wb.create_sheet("Matrix")
    matrix["A1"] = "Final PnL Matrix"
    matrix["A1"].font = Font(bold=True, size=14)
    row_no = 3
    for start_time in START_TIMES:
        matrix.cell(row_no, 1).value = f"Start {start_time}"
        matrix.cell(row_no, 1).font = Font(bold=True)
        row_no += 1
        for col, header in enumerate(("Threshold", "No Restrike", "Restrike"), start=1):
            cell = matrix.cell(row_no, col)
            cell.value = header
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        row_no += 1
        for threshold in THRESHOLDS:
            matrix.cell(row_no, 1).value = threshold
            for col, restrike in ((2, "No"), (3, "Yes")):
                match = next(
                    item
                    for item in results
                    if item["start_time"] == start_time
                    and item["threshold"] == threshold
                    and item["restrike"] == restrike
                )
                matrix.cell(row_no, col).value = match["final_pnl"]
                matrix.cell(row_no, col).number_format = "#,##0.00"
            row_no += 1
        row_no += 2

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A7" if sheet.title == "Scenario PnL" else "A1"
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 18
    wb.save(output)


if __name__ == "__main__":
    main()
