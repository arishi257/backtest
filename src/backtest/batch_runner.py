from __future__ import annotations

import argparse
from dataclasses import dataclass
from datetime import date, datetime, timedelta
import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
import pandas as pd

from backtest.config import (
    DEFAULT_PROCESSED_OUTPUT_DIR,
    DEFAULT_UNDERLYING,
    DEFAULT_WORKBOOK,
    normalize_file_date,
    normalize_underlying,
    resolve_sample_csv,
)
from backtest.data import load_option_dataset
from backtest.data import parse_expiry_text, ticker_pattern
from backtest.headless_portfolio import (
    GammaDiffTracker,
    ParkGammaTracker,
    HeadlessFrozenIvState,
    HeadlessPortfolioState,
    is_top_move_time,
)
from backtest.processed_data import ProcessedDataWriter
from backtest.replay import CsvReplayFeed, register_token_tickers
from backtest.sessions import build_backtest_sessions
from backtest.spot_data import load_spot_series


@dataclass
class BatchResult:
    date_key: str
    underlying: str
    status: str
    detail: str
    portfolio_total_pnl: float | None = None
    portfolio_gamma_diff_total: float | None = None
    park_gamma_pnl_diff_total: float | None = None
    gk_gamma_pnl_diff_total: float | None = None
    frozen_iv_total_pnl: float | None = None
    close_to_close_vol: float | None = None
    park_vol: float | None = None
    gk_vol: float | None = None


@dataclass
class VolMetrics:
    close_to_close_vol: float | None = None
    park_vol: float | None = None
    gk_vol: float | None = None


class OhlcVolTracker:
    def __init__(self) -> None:
        self.previous_bar = None
        self.close_to_close_variance = 0.0
        self.park_variance = 0.0
        self.gk_variance = 0.0
        self.has_values = False
        self.calendar_days: float | None = None
        self.intraday_var: float | None = None

    def update(
        self,
        timestamp: datetime,
        bar,
        calendar_days: float,
        intraday_var: float,
    ) -> None:
        self.calendar_days = calendar_days
        self.intraday_var = intraday_var
        if bar is None:
            return
        if self.previous_bar is None:
            self.previous_bar = bar
            return
        if is_top_move_time(timestamp):
            self.close_to_close_variance += close_to_close_variance(
                self.previous_bar.close,
                bar.close,
            )
            self.park_variance += parkinson_variance(bar.high, bar.low)
            self.gk_variance += garman_klass_variance(
                bar.open,
                bar.high,
                bar.low,
                bar.close,
            )
            self.has_values = True
        self.previous_bar = bar

    def metrics(self) -> VolMetrics:
        if (
            not self.has_values
            or self.calendar_days is None
            or self.intraday_var is None
            or self.intraday_var <= 0
        ):
            return VolMetrics()
        return VolMetrics(
            close_to_close_vol=scaled_volatility(
                self.close_to_close_variance,
                self.calendar_days,
                self.intraday_var,
            ),
            park_vol=scaled_volatility(
                self.park_variance,
                self.calendar_days,
                self.intraday_var,
            ),
            gk_vol=scaled_volatility(
                self.gk_variance,
                self.calendar_days,
                self.intraday_var,
            ),
        )


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Run headless backtests for many dates and save processed_data CSVs."
    )
    parser.add_argument("--year", help="Year to process, e.g. 2026.")
    parser.add_argument("--start-date", help="Start date, e.g. 01012026.")
    parser.add_argument("--end-date", help="End date, e.g. 31012026.")
    parser.add_argument("--dates", help="Comma-separated dates, e.g. 01012026,06012026.")
    parser.add_argument(
        "--underlying",
        default=DEFAULT_UNDERLYING,
        help="Underlying or comma-separated underlyings, e.g. NIFTY,SENSEX.",
    )
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument(
        "--processed-output-dir",
        type=Path,
        default=DEFAULT_PROCESSED_OUTPUT_DIR,
    )
    parser.add_argument(
        "--hedge-threshold",
        type=float,
        default=1.3,
        help="BS delta lots threshold for re-hedging.",
    )
    parser.add_argument(
        "--refresh-ms",
        type=positive_int,
        default=None,
        help="Milliseconds between replay slices, e.g. 20.",
    )
    parser.add_argument(
        "--excel-output",
        type=Path,
        default=None,
        help="Optional .xlsx path for OK 0DTE batch results only.",
    )
    parser.add_argument(
        "--all-dates",
        action="store_true",
        help="Run every requested date instead of only 0DTE expiry dates.",
    )
    args = parser.parse_args()

    date_keys = requested_dates(args)
    underlyings = parse_underlyings(args.underlying)
    results = []
    for date_key in date_keys:
        for underlying in underlyings:
            print(f"Running {underlying} {date_key}...")
            result = run_one(
                date_key,
                underlying,
                args.workbook,
                args.processed_output_dir,
                args.hedge_threshold,
                args.refresh_ms,
                expiry_only=not args.all_dates,
            )
            results.append(result)
            print(f"  {result.status}: {result.detail}")

    ok = sum(1 for result in results if result.status == "OK")
    skipped = len(results) - ok
    print(f"Batch complete. OK: {ok}. Skipped/failed: {skipped}.")
    print_results_table(results)
    if args.excel_output is not None:
        write_results_excel(results, args.excel_output)
        print(f"Excel output: {args.excel_output}")


def run_one(
    date_key: str,
    underlying: str,
    workbook: Path,
    processed_output_dir: Path,
    hedge_threshold: float,
    refresh_ms: int | None,
    expiry_only: bool = True,
) -> BatchResult:
    try:
        csv_path = resolve_sample_csv(date_key, underlying=underlying)
        if expiry_only:
            expiry_status = raw_csv_0dte_status(csv_path, date_key, underlying)
            if expiry_status is not None:
                return BatchResult(date_key, underlying, "SKIP", expiry_status)
        dataset = load_option_dataset(csv_path, underlying)
    except Exception as exc:
        return BatchResult(date_key, underlying, "SKIP", str(exc))
    future_series = dataset.future_series
    if future_series is None:
        print("  Nearest futures data unavailable.")

    spot_points = []
    try:
        spot_points = load_spot_series(dataset.trade_date, dataset.underlying).points
    except (FileNotFoundError, OSError, ValueError) as exc:
        print(f"  Spot data unavailable: {exc}")

    current_time = lambda: replay.now()
    try:
        sessions = build_backtest_sessions(dataset, workbook, current_time, refresh_ms=refresh_ms)
    except Exception as exc:
        return BatchResult(date_key, underlying, "SKIP", str(exc))

    replay = CsvReplayFeed(dataset, sessions)
    register_token_tickers(sessions, dataset)
    processed_writer = ProcessedDataWriter(processed_output_dir, sessions)
    universal_mid_points = {session.spec.tab_name: [] for session in sessions}
    portfolio_states = {
        session.spec.tab_name: HeadlessPortfolioState(session=session)
        for session in sessions
    }
    frozen_states = {
        session.spec.tab_name: HeadlessFrozenIvState(session)
        for session in sessions
    }
    gamma_trackers = {session.spec.tab_name: GammaDiffTracker() for session in sessions}
    park_gamma_trackers = {
        session.spec.tab_name: ParkGammaTracker() for session in sessions
    }
    vol_trackers = {session.spec.tab_name: OhlcVolTracker() for session in sessions}
    final_metrics: dict[str, float | None] = {}

    cycles = 0
    analytics = 0
    try:
        while replay.advance():
            cycles += 1
            for session in sessions:
                result = session.analytics.calculate(session.store.snapshot())
                if result is None:
                    continue
                analytics += 1
                tab_name = session.spec.tab_name
                timestamp = replay.now()
                universal_mid_points[tab_name].append((timestamp, result.universal_mid))
                portfolio_metrics = portfolio_states[tab_name].update(
                    result,
                    timestamp,
                    session.config.market.funding_rate,
                    session.config.market.brokerage_rate,
                    hedge_threshold,
                )
                frozen_metrics = frozen_states[tab_name].update(
                    result,
                    timestamp,
                    session.config.market.funding_rate,
                    session.config.market.brokerage_rate,
                    hedge_threshold,
                )
                gamma_diff_total = gamma_trackers[tab_name].update(
                    timestamp,
                    result.universal_mid,
                    portfolio_metrics.gamma_l,
                )
                park_gamma_metrics = park_gamma_trackers[tab_name].update(
                    future_series.bar_at(timestamp) if future_series else None,
                    portfolio_metrics.gamma_l,
                )
                vol_trackers[tab_name].update(
                    timestamp,
                    future_series.bar_at(timestamp) if future_series else None,
                    session.config.market.calendar_days,
                    result.intraday_var,
                )
                vol_metrics = vol_trackers[tab_name].metrics()
                final_metrics = {
                    "portfolio_total_pnl": portfolio_metrics.total_pnl,
                    "portfolio_gamma_diff_total": gamma_diff_total,
                    "park_gamma_pnl_diff_total": park_gamma_metrics.park_gamma_pnl_diff_total,
                    "gk_gamma_pnl_diff_total": park_gamma_metrics.gk_gamma_pnl_diff_total,
                    "frozen_iv_total_pnl": frozen_metrics.total_pnl,
                    "close_to_close_vol": vol_metrics.close_to_close_vol,
                    "park_vol": vol_metrics.park_vol,
                    "gk_vol": vol_metrics.gk_vol,
                }
                processed_writer.write(
                    timestamp,
                    session,
                    result,
                    spot_points,
                    universal_mid_points[tab_name],
                    portfolio_metrics.total_pnl,
                    portfolio_metrics.gamma_l,
                    gamma_diff_total,
                    frozen_metrics.total_pnl,
                    park_gamma_metrics,
                )
    except Exception as exc:
        return BatchResult(date_key, underlying, "FAIL", str(exc))

    return BatchResult(
        date_key,
        underlying,
        "OK",
        f"{cycles} cycles, {analytics} analytics rows",
        portfolio_total_pnl=final_metrics.get("portfolio_total_pnl"),
        portfolio_gamma_diff_total=final_metrics.get("portfolio_gamma_diff_total"),
        park_gamma_pnl_diff_total=final_metrics.get("park_gamma_pnl_diff_total"),
        gk_gamma_pnl_diff_total=final_metrics.get("gk_gamma_pnl_diff_total"),
        frozen_iv_total_pnl=final_metrics.get("frozen_iv_total_pnl"),
        close_to_close_vol=final_metrics.get("close_to_close_vol"),
        park_vol=final_metrics.get("park_vol"),
        gk_vol=final_metrics.get("gk_vol"),
    )


def positive_int(value: str) -> int:
    parsed = int(value)
    if parsed <= 0:
        raise argparse.ArgumentTypeError("value must be greater than zero.")
    return parsed


def print_results_table(results: list[BatchResult]) -> None:
    columns = result_columns(include_detail=True)
    rows = [
        [format_table_value(attr, getattr(result, attr)) for attr, _ in columns]
        for result in results
    ]
    widths = [
        max(len(header), *(len(row[index]) for row in rows))
        for index, (_, header) in enumerate(columns)
    ]
    header = " | ".join(
        label.ljust(widths[index])
        for index, (_, label) in enumerate(columns)
    )
    separator = "-+-".join("-" * width for width in widths)
    print(header)
    print(separator)
    for row in rows:
        print(" | ".join(value.ljust(widths[index]) for index, value in enumerate(row)))


def result_columns(include_detail: bool) -> list[tuple[str, str]]:
    columns = [
        ("date_key", "date"),
        ("underlying", "underlying"),
        ("status", "status"),
        ("portfolio_total_pnl", "portfolio_total_pnl"),
        ("portfolio_gamma_diff_total", "portfolio_gamma_diff_total"),
        ("park_gamma_pnl_diff_total", "park_gamma_pnl_diff_total"),
        ("gk_gamma_pnl_diff_total", "gk_gamma_pnl_diff_total"),
        ("frozen_iv_total_pnl", "frozen_iv_total_pnl"),
        ("close_to_close_vol", "close_to_close_vol"),
        ("park_vol", "park_vol"),
        ("gk_vol", "gk_vol"),
    ]
    if include_detail:
        columns.append(("detail", "detail"))
    return columns


def write_results_excel(results: list[BatchResult], output_path: Path) -> None:
    ok_results = [result for result in results if result.status == "OK"]
    columns = result_columns(include_detail=False)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "0DTE Results"
    headers = [label for _, label in columns]
    sheet.append(headers)
    for result in ok_results:
        sheet.append([excel_value(attr, getattr(result, attr)) for attr, _ in columns])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"

    if ok_results:
        table = Table(displayName="Sensex0DteResults", ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)

    vol_fields = {"close_to_close_vol", "park_vol", "gk_vol"}
    pnl_fields = {
        "portfolio_total_pnl",
        "portfolio_gamma_diff_total",
        "park_gamma_pnl_diff_total",
        "gk_gamma_pnl_diff_total",
        "frozen_iv_total_pnl",
    }
    for column_index, (attr, _) in enumerate(columns, start=1):
        letter = get_column_letter(column_index)
        for row_index in range(2, sheet.max_row + 1):
            cell = sheet[f"{letter}{row_index}"]
            if attr in vol_fields:
                cell.number_format = "0.00%"
            elif attr in pnl_fields:
                cell.number_format = "#,##0"

    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
        sheet.column_dimensions[letter].width = min(max(max_length + 2, 11), 28)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def excel_value(attr: str, value: object) -> object:
    if value is None:
        return None
    if attr == "date_key":
        return datetime.strptime(str(value), "%d%m%Y").date()
    return value


def format_table_value(attr: str, value: object) -> str:
    if value is None:
        return "--"
    if isinstance(value, float):
        if attr in {"close_to_close_vol", "park_vol", "gk_vol"}:
            return f"{value:.2%}"
        return f"{value:.0f}"
    return str(value)


def close_to_close_variance(previous_close: float, close: float) -> float:
    if previous_close <= 0 or close <= 0:
        return 0.0
    return math.log(close / previous_close) ** 2


def parkinson_variance(high: float, low: float) -> float:
    if high <= 0 or low <= 0:
        return 0.0
    return math.log(high / low) ** 2 / (4 * math.log(2))


def garman_klass_variance(
    open_price: float,
    high: float,
    low: float,
    close: float,
) -> float:
    if open_price <= 0 or high <= 0 or low <= 0 or close <= 0:
        return 0.0
    range_term = 0.5 * math.log(high / low) ** 2
    close_open_term = (2 * math.log(2) - 1) * math.log(close / open_price) ** 2
    return max(range_term - close_open_term, 0.0)


def scaled_volatility(
    variance: float,
    calendar_days: float,
    intraday_var: float,
) -> float:
    return math.sqrt(variance / intraday_var * calendar_days)


def requested_dates(args: argparse.Namespace) -> list[str]:
    if args.dates:
        return [normalize_file_date(value) for value in args.dates.split(",") if value.strip()]
    if args.year:
        start = date(int(args.year), 1, 1)
        end = date(int(args.year), 12, 31)
        return weekday_date_keys(start, end)
    if args.start_date and args.end_date:
        start = datetime.strptime(normalize_file_date(args.start_date), "%d%m%Y").date()
        end = datetime.strptime(normalize_file_date(args.end_date), "%d%m%Y").date()
        return weekday_date_keys(start, end)
    raise SystemExit("Use --year, --dates, or --start-date with --end-date.")


def weekday_date_keys(start: date, end: date) -> list[str]:
    if end < start:
        raise SystemExit("--end-date must be on or after --start-date.")
    current = start
    values = []
    while current <= end:
        if current.weekday() < 5:
            values.append(current.strftime("%d%m%Y"))
        current += timedelta(days=1)
    return values


def parse_underlyings(value: str) -> list[str]:
    return [normalize_underlying(part) for part in value.split(",") if part.strip()]


def raw_csv_0dte_status(csv_path: Path, date_key: str, underlying: str) -> str | None:
    trade_date = datetime.strptime(date_key, "%d%m%Y").date()
    raw = pd.read_csv(csv_path, usecols=["Ticker"])
    parsed = raw["Ticker"].astype(str).str.extract(ticker_pattern(underlying))
    expiry_text = parsed.loc[parsed["underlying"].eq(underlying), "expiry_text"].dropna()
    if expiry_text.empty:
        return "no matching option tickers"
    expiries = sorted(parse_expiry_text(expiry_text, underlying).dt.date.unique())
    if not expiries:
        return "no expiries found"
    nearest_expiry = expiries[0]
    if nearest_expiry != trade_date:
        return f"not 0DTE; nearest expiry is {nearest_expiry.isoformat()}"
    return None


if __name__ == "__main__":
    main()
